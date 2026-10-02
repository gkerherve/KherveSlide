"""Export a presentation to PowerPoint (.pptx).

Two ways, chosen in the export dialog:

* **Editable** (the default) — what PowerPoint can draw becomes real,
  editable PowerPoint content: the frame title is the slide's title
  placeholder (so PowerPoint's outline and accessibility see it), text
  boxes keep their words, bold / italic / colours, bullets and numbering,
  pictures are the original files (cropped, rotated and faded natively),
  tables are PowerPoint tables, shapes and lines are PowerPoint shapes and
  freeforms, videos are embedded. The beamer theme itself — title bar,
  head / foot lines, slide numbers, background and the master — is each
  slide's background picture, compiled with the titles left blank so the
  editable title sits exactly on it. What PowerPoint cannot draw (LaTeX
  equations, beamer blocks and theorems, pictures with effects) goes in as
  a picture cut from the compiled PDF, at its place — pixel-exact.
* **Exact look** — every slide is one picture of the compiled PDF page:
  identical to the PDF, not editable.

Hidden slides are exported hidden. Everything is measured from the same
compiled LaTeX the PDF comes from, so the PowerPoint lines up with it.
"""
from __future__ import annotations

import copy
import io
import math
import re
import tempfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

MODE_EDITABLE, MODE_EXACT = "editable", "exact"

#: KherveSlide's theme typefaces (serializer.FONT_FAMILIES keys) → the
#: font PowerPoint uses for them.
PPT_FONTS = {"helvetica": "Arial", "fira": "Fira Sans",
             "sourcesans": "Source Sans Pro", "lato": "Lato",
             "opensans": "Open Sans", "roboto": "Roboto",
             "carlito": "Calibri", "times": "Times New Roman",
             "palatino": "Palatino Linotype", "charter": "Charter"}
#: Beamer's own Latin Modern faces, per box family, as PowerPoint fonts.
#: Calibri is about as wide as Latin Modern Sans (Arial is ~10 % wider),
#: so lines break where they do in the PDF.
DEFAULT_FONTS = {"sf": "Calibri", "rm": "Cambria", "tt": "Consolas"}

_RASTER = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff")
# LaTeX PowerPoint cannot show as text: those boxes become exact pictures.
_COMPLEX = re.compile(
    r"\\(?:d?t?frac|sqrt|i*int|oint|sum|prod|lim|binom|over(?:line|brace)|"
    r"under(?:line|brace)|hat|widehat|vec|bar|dot|ddot|tilde|widetilde|"
    r"ce\b|chemfig|includegraphics|tikz|begin\{(?!itemize|enumerate)|"
    r"left|right|mathbb|mathcal|matrix|pmatrix|bmatrix|cases)|\\\[|\$\$")


class ExportError(RuntimeError):
    """The presentation could not be compiled for the export."""


@dataclass
class ExportOptions:
    mode: str = MODE_EDITABLE
    include_hidden: bool = True
    dpi: int = 200                  # background / picture resolution
    accent: str = "#1F4E79"         # bullet colour (the theme's)


@dataclass
class ExportReport:
    slides: int = 0
    editable: int = 0               # objects exported as PowerPoint objects
    pictures: int = 0               # objects exported as exact pictures
    notes: list = field(default_factory=list)


# ------------------------------------------------------------ text → runs
class _Paras(HTMLParser):
    """Paragraphs and runs from canvas.latex_to_html(pretty=True)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.paras: list[dict] = []
        self._lists: list[str] = []
        self._fmt: list[dict] = [{}]
        self._cur: dict | None = None

    def _para(self, kind="p", level=0):
        self._cur = {"kind": kind, "level": level, "runs": []}
        self.paras.append(self._cur)

    def handle_starttag(self, tag, attrs):
        fmt = dict(self._fmt[-1])
        if tag == "div":
            self._para()
        elif tag == "br":
            self._para()
            self._cur = None
            return
        elif tag in ("ul", "ol"):
            self._lists.append(tag)
        elif tag == "li":
            self._para(self._lists[-1] if self._lists else "ul",
                       max(0, len(self._lists) - 1))
        elif tag in ("b", "strong"):
            fmt["bold"] = True
        elif tag in ("i", "em"):
            fmt["italic"] = True
        elif tag == "u":
            fmt["underline"] = True
        elif tag in ("sup", "sub"):
            fmt["baseline"] = 30000 if tag == "sup" else -25000
        elif tag == "span":
            m = re.search(r"color:\s*(#[0-9A-Fa-f]{6})",
                          dict(attrs).get("style", "") or "")
            if m:
                fmt["color"] = m.group(1)
        if tag not in ("br",):
            self._fmt.append(fmt)

    def handle_endtag(self, tag):
        if tag in ("ul", "ol") and self._lists:
            self._lists.pop()
        if tag != "br" and len(self._fmt) > 1:
            self._fmt.pop()
        if tag in ("div", "li"):
            self._cur = None

    def handle_data(self, data):
        if not data:
            return
        if self._cur is None:
            if not data.strip():
                return
            kind = self._lists[-1] if self._lists else "p"
            self._para(kind, max(0, len(self._lists) - 1))
            if kind != "p":            # text after a nested list: no bullet
                self._cur["kind"] = "cont"
        self._cur["runs"].append((data, dict(self._fmt[-1])))


def text_paragraphs(latex: str) -> list[dict]:
    """A box's LaTeX as paragraphs — {"kind": p | ul | ol | cont,
    "level", "runs": [(text, {bold, italic, underline, baseline,
    color})]} — the way the canvas shows it."""
    from .canvas import latex_to_html
    parser = _Paras()
    parser.feed(latex_to_html(latex or "", pretty=True))
    parser.close()
    paras = parser.paras
    while paras and not any(t.strip() for t, _f in paras[-1]["runs"]):
        paras.pop()
    return paras or [{"kind": "p", "level": 0, "runs": []}]


def needs_picture(latex: str) -> bool:
    """True when the text holds LaTeX PowerPoint cannot show as text
    (fractions, roots, integrals, environments, chemistry…)."""
    if _COMPLEX.search(latex or ""):
        return True
    plain = "".join(t for p in text_paragraphs(latex) for t, _f in p["runs"])
    return "\\" in plain


# ------------------------------------------------------------- geometry
class _Geo:
    """Slide fractions (inside the deck's gap) → page points → EMU."""

    def __init__(self, deck, page_w_pt, page_h_pt, slide_w, slide_h):
        self.g = float(getattr(deck, "gap", 0.0) or 0.0)
        self.pw, self.ph = page_w_pt, page_h_pt
        self.sw, self.sh = slide_w, slide_h
        # PowerPoint pt per PDF pt, for font sizes
        self.font = (slide_h / 12700.0) / page_h_pt

    def frac_rect(self, o):
        g = self.g
        return (g + o.x * (1 - 2 * g), g + o.y * (1 - 2 * g),
                o.w * (1 - 2 * g), o.h * (1 - 2 * g))

    def emu(self, o):
        fx, fy, fw, fh = self.frac_rect(o)
        return (int(fx * self.sw), int(fy * self.sh),
                max(1, int(fw * self.sw)), max(1, int(fh * self.sh)))

    def page_rect(self, o):
        fx, fy, fw, fh = self.frac_rect(o)
        return (fx * self.pw, fy * self.ph, (fx + fw) * self.pw,
                (fy + fh) * self.ph)

    def pt_to_emu(self, x, y):
        return int(x / self.pw * self.sw), int(y / self.ph * self.sh)


# ------------------------------------------------------- PDF rendering
def _png(doc, index: int, dpi: int, clip=None) -> bytes:
    import pymupdf
    page = doc[index]
    zoom = dpi / 72.0
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False,
                          clip=pymupdf.Rect(*clip) if clip else None)
    return pix.tobytes("png")


def _compile(tex: str, name: str, source_dir, compile_fn):
    import pymupdf
    if compile_fn is None:
        from .compiler import compile_tex as compile_fn
    workdir = Path(tempfile.gettempdir()) / "kherveslide_pptx" / name
    res = compile_fn(tex, workdir, basename=name, source_dir=source_dir)
    pdf = getattr(res, "pdf_path", None)
    if not getattr(res, "ok", False) or not pdf or not Path(pdf).exists():
        log = (getattr(res, "log", "") or "").strip().splitlines()[-12:]
        raise ExportError("The presentation could not be compiled:\n"
                          + "\n".join(log))
    return pymupdf.open(str(pdf))


def _spans(page) -> list[tuple]:
    out = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for s in line.get("spans", []):
                if s["text"].strip():
                    out.append((s["text"], tuple(round(v, 1) for v in
                                                 s["bbox"]),
                                s["size"], s["color"], s["flags"],
                                s["font"]))
    return out


def title_spans(with_title, without_title) -> list[tuple]:
    """The frame title's text spans: on the page compiled with the titles
    and not on the one compiled with them blanked out."""
    blank = {(s[0], s[1]) for s in _spans(without_title)}
    return [s for s in _spans(with_title) if (s[0], s[1]) not in blank]


# ------------------------------------------------------------- XML bits
def _rgb(hex_colour: str):
    from pptx.dml.color import RGBColor
    h = (hex_colour or "#000000").lstrip("#")
    return RGBColor.from_string(h.upper() if len(h) == 6 else "000000")


def _set_alpha(fill_or_line_el, opacity: float) -> None:
    """Give the srgbClr inside *el* an alpha (PowerPoint transparency)."""
    if opacity >= 0.999:
        return
    from pptx.oxml.ns import qn
    for clr in fill_or_line_el.iter(qn("a:srgbClr")):
        for old in clr.findall(qn("a:alpha")):
            clr.remove(old)
        a = clr.makeelement(qn("a:alpha"), {})
        a.set("val", str(int(max(0.0, opacity) * 100000)))
        clr.append(a)


_DASH = {"dashed": "dash", "dotted": "sysDot"}


def _style_line(line, colour: str, width_pt: float, style: str = "solid",
                opacity: float = 1.0, scale: float = 1.0) -> None:
    from pptx.util import Pt
    if not colour:
        line.fill.background()
        return
    line.color.rgb = _rgb(colour)
    line.width = Pt(max(0.25, width_pt * scale))
    if style in _DASH:
        from pptx.oxml.ns import qn
        ln = line._get_or_add_ln()
        for old in ln.findall(qn("a:prstDash")):
            ln.remove(old)
        d = ln.makeelement(qn("a:prstDash"), {})
        d.set("val", _DASH[style])
        ln.append(d)
    _set_alpha(line._get_or_add_ln(), opacity)


def _arrowheads(line, start: bool, end: bool) -> None:
    from pptx.oxml.ns import qn
    ln = line._get_or_add_ln()
    for tag, on in (("a:headEnd", start), ("a:tailEnd", end)):
        for old in ln.findall(qn(tag)):
            ln.remove(old)
        if on:
            el = ln.makeelement(qn(tag), {})
            el.set("type", "triangle")
            el.set("w", "med")
            el.set("len", "med")
            ln.append(el)


def _fill(shape, colour: str, colour2: str = "", gradient: str = "vertical",
          opacity: float = 1.0) -> None:
    if not colour:
        shape.fill.background()
        return
    if colour2:
        shape.fill.gradient()
        shape.fill.gradient_angle = 90 if gradient == "vertical" else 0
        stops = shape.fill.gradient_stops
        stops[0].color.rgb = _rgb(colour)
        stops[0].position = 0.0
        stops[1].color.rgb = _rgb(colour2)
        stops[1].position = 1.0
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(colour)
    _set_alpha(shape._element.spPr, opacity)


def _set_background(slide, png: bytes) -> None:
    """The slide's background is *png*, stretched (not a shape anyone can
    move by accident)."""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    _part, rid = slide.part.get_or_add_image_part(io.BytesIO(png))
    bg = parse_xml(
        f"<p:bg {nsdecls('p', 'a', 'r')}><p:bgPr>"
        f"<a:blipFill dpi='0' rotWithShape='1'><a:blip r:embed='{rid}'/>"
        "<a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill>"
        "<a:effectLst/></p:bgPr></p:bg>")
    c_sld = slide._element.cSld
    for old in c_sld.findall(bg.tag):
        c_sld.remove(old)
    c_sld.insert(0, bg)


#: beamer's list geometry, measured on its PDF (page points): the text
#: of a first-level item starts 21 pt in (17 pt more per level, whatever
#: the font size), a bullet 5 pt in, a number 2 pt in.
LIST_MARGIN, LIST_STEP, BULLET_AT, NUMBER_AT = 21.0, 17.0, 5.0, 2.0


_BULLET_CHARS = "▶►▸▪■◆●•‣"


def theme_bullet(doc) -> tuple[str, str] | None:
    """(glyph, #colour) of the theme's first-level bullet when the PDF
    sets it as a character (triangles, squares…); None for drawn ones
    (beamer's shaded balls)."""
    seen: dict = {}
    for page in doc:
        for text, _bbox, _size, colour, _flags, _font in _spans(page):
            t = text.strip()
            if len(t) == 1 and t in _BULLET_CHARS:
                seen[(t, colour)] = seen.get((t, colour), 0) + 1
    if not seen:
        return None
    (glyph, colour), _n = max(seen.items(), key=lambda kv: kv[1])
    return glyph, "#%06X" % colour


def _bullet(paragraph, kind: str, level: int, size_pt: float,
            accent: str, pt_scale: float = 1.0, glyph: str = "•") -> None:
    """beamer-like bullets / numbers at beamer's indents; *pt_scale* turns
    page points into PowerPoint points."""
    from pptx.oxml.ns import qn
    p_pr = paragraph._p.get_or_add_pPr()
    emu = 12700 * pt_scale
    margin = int((LIST_MARGIN + LIST_STEP * level) * emu)
    if kind in ("ul", "ol"):
        at = NUMBER_AT if kind == "ol" else BULLET_AT
        p_pr.set("marL", str(margin))
        p_pr.set("indent", str(-int((LIST_MARGIN - at) * emu)))
        clr = p_pr.makeelement(qn("a:buClr"), {})
        s = clr.makeelement(qn("a:srgbClr"), {})
        s.set("val", (accent or "#1F4E79").lstrip("#").upper())
        clr.append(s)
        p_pr.append(clr)
        if kind == "ol":
            b = p_pr.makeelement(qn("a:buAutoNum"), {})
            b.set("type", ("arabicPeriod", "alphaLcParenR",
                           "romanLcPeriod")[min(level, 2)])
        else:
            if level:                   # beamer: the same mark, smaller
                sz = p_pr.makeelement(qn("a:buSzPct"), {})
                sz.set("val", str(max(60000, 100000 - 15000 * level)))
                p_pr.append(sz)
            b = p_pr.makeelement(qn("a:buChar"), {})
            b.set("char", glyph)
        p_pr.append(b)
    else:
        p_pr.set("marL", str(margin if kind == "cont" else 0))
        p_pr.set("indent", "0")
        p_pr.append(p_pr.makeelement(qn("a:buNone"), {}))


def fill_text_frame(tf, paras: list[dict], *, size_pt: float, font: str,
                    colour: str, align: str = "left", bold=False,
                    italic=False, accent: str = "#1F4E79",
                    pt_scale: float = 1.0, glyph: str = "•") -> None:
    """Write *paras* (text_paragraphs) into a PowerPoint text frame."""
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Emu, Pt
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    al = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
          "justify": PP_ALIGN.JUSTIFY}.get(align, PP_ALIGN.LEFT)
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = al
        p.line_spacing = Pt(size_pt * 1.2)        # TeX's baselineskip
        if para["kind"] in ("ul", "ol") and k:
            p.space_before = Pt(size_pt * 0.25)
        _bullet(p, para["kind"], para["level"], size_pt, accent, pt_scale,
                glyph)
        runs = para["runs"] or [("", {})]
        for text, fmt in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.size = Pt(max(1.0, size_pt))
            f.name = font
            f.bold = bool(fmt.get("bold") or bold) or None
            f.italic = bool(fmt.get("italic") or italic) or None
            if fmt.get("underline"):
                f.underline = True
            f.color.rgb = _rgb(fmt.get("color") or colour)
            if fmt.get("baseline"):
                r._r.get_or_add_rPr().set("baseline", str(fmt["baseline"]))


# ------------------------------------------------------------- exporter
class _Exporter:
    def __init__(self, deck, deck_dir, opts: ExportOptions, compile_fn,
                 progress):
        self.deck = deck
        self.dir = Path(deck_dir) if deck_dir else None
        self.opts = opts
        self.compile_fn = compile_fn
        self.progress = progress or (lambda _t, _f: None)
        self.report = ExportReport()

    # -- helpers
    def _path(self, p: str) -> Path | None:
        if not p:
            return None
        path = Path(p)
        if not path.is_absolute() and self.dir is not None:
            path = self.dir / path
        return path if path.exists() else None

    def _font(self, family: str = "") -> str:
        from .serializer import deck_body_family, deck_typeface
        if family in ("rm", "tt"):
            return DEFAULT_FONTS[family]
        face = deck_typeface(self.deck)
        if face in PPT_FONTS and family in ("", "sf"):
            return PPT_FONTS[face]
        if family == "sf":
            return DEFAULT_FONTS["sf"]
        return DEFAULT_FONTS[deck_body_family(self.deck)]

    def _elements(self, i):
        """Everything drawn on page *i* that is not the theme: text spans,
        vector drawings and images, as rects (cached per page)."""
        cache = getattr(self, "_els", None)
        if cache is None:
            cache = self._els = {}
        if i not in cache:
            def rects(page):
                out = [tuple(s[1]) for s in _spans(page)]
                for d in page.get_drawings():
                    r = d.get("rect")
                    if r is not None:
                        out.append((r.x0, r.y0, r.x1, r.y1))
                for im in page.get_image_info():
                    out.append(tuple(im["bbox"]))
                return out
            theme = set()
            if getattr(self, "with_titles", None) is not None:
                theme = {tuple(round(v) for v in r)
                         for r in rects(self.with_titles[i])}
            cache[i] = [r for r in rects(self.full[i])
                        if tuple(round(v) for v in r) not in theme]
        return cache[i]

    def _content_rect(self, i, o):
        """The box's rect grown to what is really drawn for it — an
        equation or a text taller than its box overflows it in the PDF.
        Drawn things touching the region join it, unless they sit in
        another object's box."""
        x0, y0, x1, y1 = self.geo.page_rect(o)
        others = [self.geo.page_rect(b) for b in self.deck.slides[i].objects
                  if b is not o]

        def owned_elsewhere(r):
            cx, cy = (r[0] + r[2]) / 2, (r[1] + r[3]) / 2
            if x0 <= cx <= x1 and y0 <= cy <= y1:
                return False
            return any(a[0] <= cx <= a[2] and a[1] <= cy <= a[3]
                       for a in others)

        els = [r for r in self._elements(i) if not owned_elsewhere(r)]
        tol = 6.0               # page pt: the gap between two text lines
        grown = True
        while grown:
            grown = False
            for r in els:
                if (r[0] <= x1 + tol and r[2] >= x0 - tol
                        and r[1] <= y1 + tol and r[3] >= y0 - tol):
                    nx0, ny0 = min(x0, r[0]), min(y0, r[1])
                    nx1, ny1 = max(x1, r[2]), max(y1, r[3])
                    if (nx0, ny0, nx1, ny1) != (x0, y0, x1, y1):
                        x0, y0, x1, y1 = nx0, ny0, nx1, ny1
                        grown = True
        return x0 - 0.5, y0 - 0.5, x1 + 0.5, y1 + 0.5

    def _crop(self, slide, i, o) -> None:
        """An object as a picture cut from the compiled page — the whole of
        what is drawn for it."""
        x0, y0, x1, y1 = self._content_rect(i, o)
        page = self.full[i].rect
        x0, y0 = max(0, x0), max(0, y0)
        x1, y1 = min(page.width, x1), min(page.height, y1)
        if x1 - x0 < 1 or y1 - y0 < 1:
            return
        png = _png(self.full, i, max(self.opts.dpi, 220), (x0, y0, x1, y1))
        left, top = self.geo.pt_to_emu(x0, y0)
        right, bottom = self.geo.pt_to_emu(x1, y1)
        pic = slide.shapes.add_picture(io.BytesIO(png), left, top,
                                       right - left, bottom - top)
        pic.name = f"{type(o).__name__[5:]} (as in the PDF)"
        self.report.pictures += 1

    # -- objects
    def _text(self, slide, i, o) -> None:
        from pptx.enum.shapes import MSO_SHAPE
        if (o.block or needs_picture(o.text)) and not o.locked:
            self._crop(slide, i, o)
            return
        x, y, w, h = self.geo.emu(o)
        framed = bool(o.fill or o.border_color)
        if framed:
            kind = (MSO_SHAPE.ROUNDED_RECTANGLE if o.corner == "rounded"
                    else MSO_SHAPE.RECTANGLE)
            shp = slide.shapes.add_shape(kind, x, y, w, h)
            _fill(shp, o.fill, o.fill2, o.gradient, o.fill_opacity)
            if o.border_color:
                _style_line(shp.line, o.border_color, o.border_width,
                            o.border_style, scale=self.geo.font)
            else:
                shp.line.fill.background()
            shp.shadow.inherit = False
        else:
            shp = slide.shapes.add_textbox(x, y, w, h)
        fill_text_frame(shp.text_frame, text_paragraphs(o.text),
                        size_pt=o.font_pt * self.geo.font,
                        font=self._font(o.font_family), colour=o.color,
                        align=o.align, bold=o.bold, italic=o.italic,
                        accent=self.accent, pt_scale=self.geo.font,
                        glyph=self.glyph)
        if framed:
            pad = int(4 * 12700 * self.geo.font)
            tf = shp.text_frame
            tf.margin_left = tf.margin_right = pad
            tf.margin_top = tf.margin_bottom = pad
        self.report.editable += 1

    def _picture_file(self, o) -> io.BytesIO | None:
        path = self._path(o.path)
        if path is None:
            return None
        ext = path.suffix.lower()
        if ext in _RASTER:
            return io.BytesIO(path.read_bytes())
        if ext == ".pdf":
            import pymupdf
            doc = pymupdf.open(str(path))
            pix = doc[0].get_pixmap(dpi=max(self.opts.dpi, 220), alpha=True)
            return io.BytesIO(pix.tobytes("png"))
        if ext in (".wmf", ".emf"):
            from .metafile import ensure_raster
            png = ensure_raster(str(path))
            return io.BytesIO(Path(png).read_bytes()) if png else None
        if ext == ".svg":
            from PySide6.QtCore import QBuffer, QByteArray, QIODevice
            from PySide6.QtGui import QImage, QPainter
            from PySide6.QtSvg import QSvgRenderer
            r = QSvgRenderer(str(path))
            size = r.defaultSize()
            k = max(1.0, 1600 / max(1, size.width()))
            img = QImage(int(size.width() * k), int(size.height() * k),
                         QImage.Format_ARGB32)
            img.fill(0)
            p = QPainter(img)
            r.render(p)
            p.end()
            ba = QByteArray()
            buf = QBuffer(ba)
            buf.open(QIODevice.WriteOnly)
            img.save(buf, "PNG")
            return io.BytesIO(bytes(ba))
        return None

    def _picture(self, slide, i, o) -> None:
        from . import image_effects
        data = None if image_effects.has_effects(o) else self._picture_file(o)
        if data is None:
            self._crop(slide, i, o)
            return
        x, y, w, h = self.geo.emu(o)
        pic = slide.shapes.add_picture(data, x, y, w, h)
        pic.crop_left, pic.crop_top = o.crop_l, o.crop_t
        pic.crop_right, pic.crop_bottom = o.crop_r, o.crop_b
        if o.rotation:
            pic.rotation = o.rotation
        if o.border_color:
            _style_line(pic.line, o.border_color, o.border_width,
                        o.border_style, scale=self.geo.font)
        if o.opacity < 0.999:
            from pptx.oxml.ns import qn
            blip = pic._element.find(".//" + qn("a:blip"))
            a = blip.makeelement(qn("a:alphaModFix"), {})
            a.set("amt", str(int(o.opacity * 100000)))
            blip.append(a)
        self.report.editable += 1

    def _table(self, slide, i, o) -> None:
        rows = [list(r) for r in (o.rows or [[""]])]
        if any(needs_picture(c) for r in rows for c in r) and not o.locked:
            self._crop(slide, i, o)
            return
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        x, y, w, h = self.geo.emu(o)
        gf = slide.shapes.add_table(len(rows), ncols, x, y, w, h)
        tbl = gf.table
        tbl.first_row = bool(o.header)
        tbl.horz_banding = False
        size = o.font_pt * self.geo.font
        for r, row in enumerate(rows):
            tbl.rows[r].height = max(1, h // len(rows))
            for c, txt in enumerate(row):
                cell = tbl.cell(r, c)
                head = o.header and r == 0
                if head and o.header_bg:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(o.header_bg)
                elif o.striped and r % 2 == (0 if o.header else 1):
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(o.stripe_color)
                elif o.fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(o.fill)
                else:
                    cell.fill.background()
                fill_text_frame(
                    cell.text_frame, text_paragraphs(txt), size_pt=size,
                    font=self._font(), colour=(o.header_fg if head
                                               else o.color),
                    align=o.align, bold=head, accent=self.accent,
                    pt_scale=self.geo.font, glyph=self.glyph)
                pad = int(3 * 12700 * self.geo.font)
                cell.margin_left = cell.margin_right = pad
                cell.margin_top = cell.margin_bottom = pad // 2
                _cell_borders(cell, r, c, len(rows), ncols, o,
                              self.geo.font)
        for c in range(ncols):
            tbl.columns[c].width = max(1, w // ncols)
        self.report.editable += 1

    def _line(self, slide, i, o) -> None:
        x, y, w, h = self.geo.emu(o)
        if o.curve:
            pts = _curve_points(o.curve)
            fb = slide.shapes.build_freeform(x, y, scale=1.0)
            fb.add_line_segments([(x + u * w, y + v * h) for u, v in pts],
                                 close=False)
            shp = fb.convert_to_shape()
            shp.fill.background()
        else:
            from pptx.enum.shapes import MSO_CONNECTOR
            shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y,
                                             x + w, y + h)
        _style_line(shp.line, o.color, o.width_pt, o.style, o.opacity,
                    scale=self.geo.font)
        _arrowheads(shp.line, o.arrow_start, o.arrow_end)
        self.report.editable += 1

    def _shape(self, slide, i, o) -> None:
        from pptx.enum.shapes import MSO_SHAPE
        from . import shapes as kshapes
        x, y, w, h = self.geo.emu(o)
        kind = kshapes.outline(o.shape)
        if kind[0] == "ellipse":
            shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
        elif kind[0] == "rect":
            rounded = kind[1] or o.corner == "rounded"
            shp = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE if rounded
                else MSO_SHAPE.RECTANGLE, x, y, w, h)
        else:
            pts = kind[1]
            fb = slide.shapes.build_freeform(x + pts[0][0] * w,
                                             y + pts[0][1] * h, scale=1.0)
            fb.add_line_segments([(x + u * w, y + v * h) for u, v in pts[1:]],
                                 close=True)
            shp = fb.convert_to_shape()
        _fill(shp, o.fill, o.fill2, o.gradient, o.opacity)
        if o.border_color:
            _style_line(shp.line, o.border_color, o.border_width, o.style,
                        o.opacity, scale=self.geo.font)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        if o.rotation:
            shp.rotation = o.rotation
        self.report.editable += 1

    def _video(self, slide, i, o) -> None:
        path = self._path(o.path)
        if path is None:
            self._crop(slide, i, o)
            return
        import mimetypes
        x, y, w, h = self.geo.emu(o)
        poster = self._path(o.poster)
        slide.shapes.add_movie(
            str(path), x, y, w, h,
            poster_frame_image=str(poster) if poster else None,
            mime_type=mimetypes.guess_type(str(path))[0] or "video/mp4")
        self.report.editable += 1

    def _title(self, slide, i) -> bool:
        """Put the frame title in the title placeholder, where and as the
        theme draws it. False when the theme draws no title here."""
        spans = title_spans(self.with_titles[i], self.no_titles[i])
        ph = slide.shapes.title
        if not spans or ph is None:
            return False
        x0 = min(s[1][0] for s in spans)
        y0 = min(s[1][1] for s in spans)
        x1 = max(s[1][2] for s in spans)
        y1 = max(s[1][3] for s in spans)
        big = max(spans, key=lambda s: s[2])
        left, top = self.geo.pt_to_emu(x0, y0)
        right, bottom = self.geo.pt_to_emu(x1, y1)
        ph.left, ph.top = left, top
        # room to edit: to the slide's right margin, a little taller
        ph.width = max(right - left, self.geo.sw - left - left // 2)
        ph.height = int((bottom - top) * 1.15)
        bold = bool(big[4] & 16) or "bold" in big[5].lower() \
            or re.search(r"(bx|Bold|-B)", big[5]) is not None
        colour = "#%06X" % big[3]
        fill_text_frame(ph.text_frame,
                        text_paragraphs(self.deck.slides[i].title),
                        size_pt=big[2] * self.geo.font, font=self._font(),
                        colour=colour, bold=bold, accent=self.accent)
        return True

    # -- the whole thing
    def run(self, out_path) -> ExportReport:
        from pptx import Presentation
        from pptx.util import Emu
        from .serializer import serialize_backdrop, serialize_deck
        deck = copy.deepcopy(self.deck)
        shown = [s for s in deck.slides
                 if self.opts.include_hidden or not s.hidden]
        if not shown:
            raise ExportError("Every slide is hidden — nothing to export.")
        hidden = [s.hidden for s in shown]
        deck.slides = shown
        for s in deck.slides:
            s.hidden = False
        self.deck = deck
        src = self.dir
        self.progress("Compiling the slides…", 0.05)
        self.full = _compile(serialize_deck(deck), "slides", src,
                             self.compile_fn)
        editable = self.opts.mode == MODE_EDITABLE
        if editable:
            self.progress("Compiling the theme backgrounds…", 0.3)
            self.with_titles = _compile(serialize_backdrop(deck), "theme",
                                        src, self.compile_fn)
            blank = copy.deepcopy(deck)
            for s in blank.slides:
                if s.title:
                    s.title = f"\\phantom{{{s.title}}}"
            self.no_titles = _compile(serialize_backdrop(blank),
                                      "theme_blank", src, self.compile_fn)
        self.glyph, self.accent = "•", self.opts.accent
        found = theme_bullet(self.full)
        if found:
            self.glyph, self.accent = found
        page = self.full[0].rect
        prs = Presentation()
        ratio = page.width / page.height
        prs.slide_height = Emu(6858000)                  # 7.5 in
        prs.slide_width = Emu(int(round(6858000 * ratio)))
        self.geo = _Geo(deck, page.width, page.height, prs.slide_width,
                        prs.slide_height)
        prs.core_properties.title = deck.title or ""
        prs.core_properties.author = deck.author or ""
        n = len(deck.slides)
        for i, s in enumerate(deck.slides):
            self.progress(f"Slide {i + 1} of {n}…", 0.45 + 0.5 * i / n)
            layout = prs.slide_layouts[5 if (editable and s.title) else 6]
            slide = prs.slides.add_slide(layout)
            if not editable:
                _set_background(slide, _png(self.full, i, self.opts.dpi))
            else:
                _set_background(slide,
                                _png(self.no_titles, i, self.opts.dpi))
                if s.title and not self._title(slide, i):
                    ph = slide.shapes.title
                    ph._element.getparent().remove(ph._element)
                for o in s.objects:
                    kind = type(o).__name__
                    handler = {"SlideText": self._text,
                               "SlidePicture": self._picture,
                               "SlideTable": self._table,
                               "SlideLine": self._line,
                               "SlideShape": self._shape,
                               "SlideVideo": self._video}.get(kind)
                    try:
                        (handler or self._crop)(slide, i, o)
                    except Exception as e:     # keep going, exactly
                        self.report.notes.append(
                            f"Slide {i + 1}: a {kind[5:].lower()} went in "
                            f"as a picture ({e}).")
                        self._crop(slide, i, o)
            if hidden[i]:
                slide._element.set("show", "0")
            self.report.slides += 1
        self.progress("Saving…", 0.97)
        prs.save(str(out_path))
        self.progress("Done", 1.0)
        return self.report


def _curve_points(curve: list, steps: int = 16) -> list[tuple]:
    """Sample SlideLine.curve (cubic segments in (u, v) box fractions,
    starting at (0, 0)) into a polyline."""
    pts = [(0.0, 0.0)]
    x0, y0 = 0.0, 0.0
    for k in range(0, len(curve) - 5, 6):
        c1x, c1y, c2x, c2y, ex, ey = curve[k:k + 6]
        for j in range(1, steps + 1):
            t = j / steps
            m = 1 - t
            pts.append((m ** 3 * x0 + 3 * m * m * t * c1x + 3 * m * t * t * c2x
                        + t ** 3 * ex,
                        m ** 3 * y0 + 3 * m * m * t * c1y + 3 * m * t * t * c2y
                        + t ** 3 * ey))
        x0, y0 = ex, ey
    return pts


def _cell_borders(cell, r, c, nrows, ncols, o, scale) -> None:
    """The table's grid (all / horizontal / outer / none) as cell
    borders."""
    from pptx.oxml.ns import qn
    grid = getattr(o, "grid", "all")
    tc_pr = cell._tc.get_or_add_tcPr()
    sides = {"a:lnL": c == 0, "a:lnR": c == ncols - 1,
             "a:lnT": r == 0, "a:lnB": r == nrows - 1}
    if grid == "all":
        sides = dict.fromkeys(sides, True)
    elif grid == "horizontal":
        sides = {"a:lnL": False, "a:lnR": False, "a:lnT": True,
                 "a:lnB": True}
    elif grid == "none":
        sides = dict.fromkeys(sides, False)
    width = str(int(max(0.25, o.rule_width * scale) * 12700))
    colour = (o.rule_color or "#000000").lstrip("#").upper()
    # tcPr wants its borders first, in the order L R T B
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
        ln = tc_pr.makeelement(qn(tag), {})
        if sides[tag]:
            ln.set("w", width)
            sf = ln.makeelement(qn("a:solidFill"), {})
            clr = sf.makeelement(qn("a:srgbClr"), {})
            clr.set("val", colour)
            sf.append(clr)
            ln.append(sf)
        else:
            ln.set("w", "0")
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tc_pr.insert(0, ln)


def export_pptx(deck, out_path, deck_dir=None, opts: ExportOptions | None
                = None, compile_fn=None, progress=None) -> ExportReport:
    """Write *deck* as a PowerPoint file at *out_path* (see the module
    docstring). *deck_dir* resolves relative picture paths; *progress* is
    called with (message, fraction)."""
    return _Exporter(deck, deck_dir, opts or ExportOptions(), compile_fn,
                     progress).run(out_path)
