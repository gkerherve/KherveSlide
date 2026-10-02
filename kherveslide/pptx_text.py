"""PowerPoint text and colour, resolved the way PowerPoint resolves them.

DrawingML text inherits almost everything: a run's size, colour, bold,
a paragraph's alignment, bullet and spacing come from the run, the
paragraph, the shape's list style, then (for placeholders) the layout and
master placeholders and the master's title/body style, else the
presentation's default text style. Colours are often theme colours
("accent1", "tx1") with luminance modifiers. This module walks that
chain and lays a text frame out into KherveSlide text boxes.

A text frame becomes one or more stacked :class:`SlideText` boxes: a new
box starts where the font size or alignment changes, or after blank
paragraphs — each placed where PowerPoint puts it (line spacing, space
before/after, insets, vertical anchor, autofit shrink), since one
KherveSlide box has a single size, alignment and fixed leading.
"""
from __future__ import annotations

import colorsys
import math
from dataclasses import dataclass, field

from .model import SlideText

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
EMU_PER_PT = 12700

_SPECIALS = [("\\", "\\textbackslash{}"), ("&", "\\&"), ("%", "\\%"),
             ("$", "\\$"), ("#", "\\#"), ("_", "\\_"), ("{", "\\{"),
             ("}", "\\}"), ("~", "\\textasciitilde{}"),
             ("^", "\\textasciicircum{}")]

# The Symbol font maps Latin letters onto Greek ("j" is phi); PowerPoint
# often stores them in the private-use area (U+F0xx).
_SYMBOL_GREEK = dict(zip(
    "abgdezhqiklmnxoprstufcywjvJ"
    "ABGDEZHQIKLMNXOPRSTUFCYWV",
    "αβγδεζηθικλμνξοπρστυφχψωϕϖϑ"
    "ΑΒΓΔΕΖΗΘΙΚΛΜΝΞΟΠΡΣΤΥΦΧΨΩς"))


def escape(text: str) -> str:
    out = text or ""
    for ch, rep in _SPECIALS:
        out = out.replace(ch, rep)
    return out


# ------------------------------------------------------------------ colours
_SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
_PRESET = {"black": "000000", "white": "FFFFFF", "red": "FF0000",
           "green": "008000", "blue": "0000FF", "yellow": "FFFF00",
           "gray": "808080", "grey": "808080", "orange": "FFA500"}


@dataclass
class Theme:
    colors: dict = field(default_factory=dict)    # dk1/lt1/accent1… → RRGGBB
    minor_font: str = ""
    major_font: str = ""

    @classmethod
    def from_presentation(cls, prs) -> "Theme":
        from lxml import etree
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        t = cls()
        try:
            part = prs.slide_masters[0].part.part_related_by(RT.THEME)
            root = etree.fromstring(part.blob)
        except Exception:
            return t
        scheme = root.find(f".//{A}clrScheme")
        if scheme is not None:
            for slot in scheme:
                name = slot.tag.replace(A, "")
                c = slot[0] if len(slot) else None
                if c is None:
                    continue
                val = c.get("lastClr") or c.get("val") or ""
                if len(val) == 6:
                    t.colors[name] = val.upper()
        for kind in ("minor", "major"):
            el = root.find(f".//{A}{kind}Font/{A}latin")
            if el is not None:
                setattr(t, f"{kind}_font", el.get("typeface", ""))
        return t


def _hsl_mod(rgb: str, mods: list[tuple[str, float]]) -> str:
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    for name, v in mods:
        if name in ("lumMod", "lumOff"):
            h, l, s = colorsys.rgb_to_hls(r, g, b)
            l = l * v if name == "lumMod" else l + v
            r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
        elif name == "shade":
            r, g, b = r * v, g * v, b * v
        elif name == "tint":
            r, g, b = (1 - (1 - c) * v for c in (r, g, b))
    return "".join(f"{int(round(max(0, min(1, c)) * 255)):02X}"
                   for c in (r, g, b))


def color_of(el, theme: Theme, ph_color: str = "") -> tuple[str, float]:
    """Resolve a colour element (srgbClr, schemeClr, sysClr, prstClr) to
    ``("#RRGGBB", alpha)``; ("", 1) when there is none."""
    if el is None:
        return "", 1.0
    tag = el.tag.replace(A, "")
    val = el.get("val") or ""
    if tag == "srgbClr":
        base = val.upper()
    elif tag == "sysClr":
        base = (el.get("lastClr") or "000000").upper()
    elif tag == "prstClr":
        base = _PRESET.get(val.lower(), "000000")
    elif tag == "schemeClr":
        if val == "phClr":
            base = ph_color.lstrip("#").upper() if ph_color else ""
        else:
            base = theme.colors.get(_SCHEME_ALIAS.get(val, val), "")
    else:
        return "", 1.0
    if len(base) != 6:
        return "", 1.0
    mods, alpha = [], 1.0
    for m in el:
        name = m.tag.replace(A, "")
        try:
            v = int(m.get("val", "0")) / 100000.0
        except ValueError:
            continue
        if name == "alpha":
            alpha = v
        elif name in ("lumMod", "lumOff", "shade", "tint"):
            mods.append((name, v))
    return "#" + (_hsl_mod(base, mods) if mods else base), alpha


def fill_color(parent, theme: Theme) -> tuple[str, float]:
    """The colour inside a ``<a:solidFill>`` under *parent*."""
    if parent is None:
        return "", 1.0
    sf = parent.find(f"{A}solidFill")
    if sf is None or not len(sf):
        return "", 1.0
    return color_of(sf[0], theme)


# --------------------------------------------------------- style inheritance
def _level_props(lst, level: int):
    if lst is None:
        return None
    return lst.find(f"{A}lvl{level + 1}pPr")


class StyleChain:
    """The list styles a shape's paragraphs inherit from, nearest first."""

    def __init__(self, shape, prs, theme: Theme):
        self.theme = theme
        self.lists = []
        txb = shape._element.find(f".//{P}txBody")
        if txb is not None:
            self.lists.append(txb.find(f"{A}lstStyle"))
        self.font_ref = ""
        style = shape._element.find(f"{P}style")
        if style is not None:
            fr = style.find(f"{A}fontRef")
            if fr is not None and len(fr):
                self.font_ref = color_of(fr[0], theme)[0]
        kind = None
        try:
            if shape.is_placeholder:
                kind = shape.placeholder_format.type
                base = getattr(shape, "_base_placeholder", None)
                while base is not None:
                    b = base._element.find(f".//{P}txBody")
                    if b is not None:
                        self.lists.append(b.find(f"{A}lstStyle"))
                    base = getattr(base, "_base_placeholder", None)
        except Exception:
            kind = None
        master = None
        try:
            master = shape.part.slide_layout.slide_master._element
        except Exception:
            pass
        if kind is not None and master is not None:
            from pptx.enum.shapes import PP_PLACEHOLDER
            name = ("titleStyle" if kind in (PP_PLACEHOLDER.TITLE,
                                             PP_PLACEHOLDER.CENTER_TITLE)
                    else "bodyStyle" if kind in (
                        PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT,
                        PP_PLACEHOLDER.SUBTITLE) else "otherStyle")
            self.lists.append(master.find(f"{P}txStyles/{P}{name}"))
        try:
            self.lists.append(prs.part._element.find(
                f"{P}defaultTextStyle"))
        except Exception:
            pass

    def para(self, level: int, getter):
        for lst in self.lists:
            lv = _level_props(lst, level)
            if lv is None:
                continue
            v = getter(lv)
            if v is not None:
                return v
        return None

    def run(self, level: int, getter):
        return self.para(level, lambda lv: (
            getter(lv.find(f"{A}defRPr"))
            if lv.find(f"{A}defRPr") is not None else None))


# --------------------------------------------------------------- paragraphs
@dataclass
class Run:
    text: str
    size: float
    bold: bool
    italic: bool
    color: str
    baseline: int = 0          # >0 superscript, <0 subscript


@dataclass
class Para:
    runs: list
    level: int
    align: str
    bullet: bool
    line_pct: float            # line spacing multiple (1.0 = single)
    line_pts: float            # exact line spacing in pt (0 = use pct)
    space_before: float        # pt
    space_after: float         # pt
    size: float                # the paragraph's dominant size (pt)
    indent: float = 0.0        # left margin in pt (marL)

    @property
    def empty(self) -> bool:
        return not "".join(r.text for r in self.runs).strip()


_ALIGN = {"l": "left", "ctr": "center", "r": "right", "just": "left",
          "dist": "center"}


def _spacing(el, size: float) -> float | None:
    """A spcBef/spcAft element → points."""
    if el is None:
        return None
    pts = el.find(f"{A}spcPts")
    if pts is not None:
        return int(pts.get("val", "0")) / 100.0
    pct = el.find(f"{A}spcPct")
    if pct is not None:
        return int(pct.get("val", "0")) / 100000.0 * size * 1.2
    return None


def _bool_attr(el, name):
    if el is None or el.get(name) is None:
        return None
    return el.get(name) in ("1", "true")


def read_paragraphs(shape, prs, theme: Theme, font_scale: float = 1.0,
                    spacing_cut: float = 0.0) -> list[Para]:
    chain = StyleChain(shape, prs, theme)
    txb = shape._element.find(f".//{P}txBody")
    if txb is None:
        return []
    paras = []
    for p in txb.findall(f"{A}p"):
        ppr = p.find(f"{A}pPr")
        level = int(ppr.get("lvl", "0")) if ppr is not None else 0

        def pget(fn, default=None):
            v = fn(ppr) if ppr is not None else None
            if v is None:
                v = chain.para(level, fn)
            return default if v is None else v

        def rget(rpr, fn, default=None):
            v = fn(rpr) if rpr is not None else None
            if v is None and ppr is not None and \
                    ppr.find(f"{A}defRPr") is not None:
                v = fn(ppr.find(f"{A}defRPr"))
            if v is None:
                v = chain.run(level, fn)
            return default if v is None else v

        align = _ALIGN.get(pget(lambda e: e.get("algn")), "left")
        bullet = pget(lambda e: (False if e.find(f"{A}buNone") is not None
                                 else True if (e.find(f"{A}buChar") is not None
                                               or e.find(f"{A}buAutoNum")
                                               is not None) else None), False)
        indent = int(pget(lambda e: e.get("marL"), "0")) / EMU_PER_PT

        def size_of(e):
            return int(e.get("sz")) / 100.0 if e is not None and e.get("sz") \
                else None

        def color_fn(e):
            if e is None:
                return None
            c = fill_color(e, theme)[0]
            return c or None

        runs = []
        for child in p:
            tag = child.tag.replace(A, "")
            if tag not in ("r", "br", "fld"):
                continue
            rpr = child.find(f"{A}rPr")
            if tag == "br":
                runs.append(Run("\n", 0, False, False, ""))
                continue
            t = child.find(f"{A}t")
            text = t.text if t is not None and t.text else ""
            font = rpr.find(f"{A}sym") if rpr is not None else None
            latin = rpr.find(f"{A}latin") if rpr is not None else None
            face = (font.get("typeface") if font is not None
                    else latin.get("typeface") if latin is not None else "")
            if face and "symbol" in face.lower():
                text = "".join(_SYMBOL_GREEK.get(
                    chr(ord(c) - 0xF000) if ord(c) >= 0xF000 else c, c)
                    for c in text)
            size = rget(rpr, size_of, 18.0) * font_scale
            colour = rget(rpr, color_fn, "")
            if rpr is not None and rpr.find(f"{A}hlinkClick") is not None \
                    and rpr.find(f"{A}solidFill") is None:
                colour = "#" + theme.colors.get("hlink", "0563C1")
            if not colour:
                colour = chain.font_ref or "#" + theme.colors.get(
                    "dk1", "000000")
            base = int(rpr.get("baseline", "0")) if rpr is not None else 0
            runs.append(Run(text, size,
                            bool(rget(rpr, lambda e: _bool_attr(e, "b"),
                                      False)),
                            bool(rget(rpr, lambda e: _bool_attr(e, "i"),
                                      False)),
                            colour.upper(), base))
        end = p.find(f"{A}endParaRPr")
        if runs:
            weights: dict = {}
            for r in runs:
                weights[r.size] = weights.get(r.size, 0) + len(r.text)
            size = max(weights, key=weights.get) if any(weights.values()) \
                else runs[0].size
        else:
            size = rget(end, size_of, 18.0) * font_scale

        ln = pget(lambda e: e.find(f"{A}lnSpc"))
        line_pct, line_pts = 1.0, 0.0
        if ln is not None:
            if ln.find(f"{A}spcPct") is not None:
                line_pct = int(ln.find(f"{A}spcPct").get("val")) / 100000.0
            elif ln.find(f"{A}spcPts") is not None:
                line_pts = int(ln.find(f"{A}spcPts").get("val")) / 100.0
        line_pct = max(0.5, line_pct - spacing_cut)
        before = pget(lambda e: _spacing(e.find(f"{A}spcBef"), size), 0.0)
        after = pget(lambda e: _spacing(e.find(f"{A}spcAft"), size), 0.0)
        paras.append(Para(runs, level, align, bool(bullet), line_pct,
                          line_pts, before, after, size, indent))
    return paras


# ------------------------------------------------------------------ layout
# Calibri's average advance is about half an em; good enough to count
# how many lines a paragraph wraps to.
_AVG_ADVANCE = 0.5


def para_lines(para: Para, width_pt: float) -> int:
    avail = max(10.0, width_pt - (para.indent if para.bullet else 0.0))
    lines, cur = 1, 0.0
    for r in para.runs:
        if r.text == "\n":
            lines += 1
            cur = 0.0
            continue
        adv = len(r.text) * r.size * _AVG_ADVANCE * (1.05 if r.bold else 1)
        cur += adv
    if cur > avail:
        lines += int(math.ceil(cur / avail)) - 1
    return lines


def line_height(para: Para) -> float:
    if para.line_pts:
        return para.line_pts
    return para.size * 1.2 * para.line_pct


def para_height(para: Para, width_pt: float) -> float:
    return para_lines(para, width_pt) * line_height(para)


def _run_latex(r: Run, box_color: str) -> str:
    if r.text == "\n":
        return ""
    t = escape(r.text)
    if not t:
        return ""
    if r.italic:
        t = f"\\textit{{{t}}}"
    if r.bold:
        t = f"\\textbf{{{t}}}"
    if r.baseline > 0:
        t = f"\\textsuperscript{{{t}}}"
    elif r.baseline < 0:
        t = f"\\textsubscript{{{t}}}"
    if r.color and r.color != box_color and t.strip():
        t = f"\\textcolor[HTML]{{{r.color.lstrip('#')}}}{{{t}}}"
    return t


def _para_lines(para: Para, box_color: str) -> list[str]:
    """A paragraph's LaTeX, split at its soft line breaks."""
    out, cur = [], []
    for r in para.runs:
        if r.text == "\n":
            out.append("".join(cur))
            cur = []
        else:
            cur.append(_run_latex(r, box_color))
    out.append("".join(cur))
    return out


def group_text(paras: list[Para], box_color: str) -> str:
    """Lines of text for one box: bulleted paragraphs become an itemize
    (nested by level); the others are plain lines."""
    lines: list[str] = []
    depth = 0
    for para in paras:
        want = (para.level + 1) if para.bullet else 0
        while depth > want:
            lines.append("\\end{itemize}")
            depth -= 1
        while depth < want:
            lines.append("\\begin{itemize}")
            depth += 1
        parts = _para_lines(para, box_color)
        if para.bullet:
            lines.append("\\item " + " \\\\ ".join(p for p in parts))
        else:
            lines.extend(parts)
    while depth:
        lines.append("\\end{itemize}")
        depth -= 1
    return "\n".join(lines)


def _dominant_color(paras: list[Para]) -> str:
    weights: dict = {}
    for p in paras:
        for r in p.runs:
            if r.text.strip():
                weights[r.color] = weights.get(r.color, 0) + len(r.text)
    return max(weights, key=weights.get) if weights else "#000000"


def text_boxes(shape, prs, theme: Theme, rect_emu, slide_w: int,
               slide_h: int) -> list[SlideText]:
    """Lay a shape's text frame out as stacked KherveSlide text boxes.
    *rect_emu* is the shape's (x, y, w, h) on the slide in EMU."""
    txb = shape._element.find(f".//{P}txBody")
    if txb is None:
        return []
    body = txb.find(f"{A}bodyPr")

    def ins(name, default):
        v = body.get(name) if body is not None else None
        return int(v) if v is not None else default

    l_in, r_in = ins("lIns", 91440), ins("rIns", 91440)
    t_in, b_in = ins("tIns", 45720), ins("bIns", 45720)
    anchor = (body.get("anchor") if body is not None else None) or "t"
    nowrap = body is not None and body.get("wrap") == "none"
    scale, cut = 1.0, 0.0
    fit = body.find(f"{A}normAutofit") if body is not None else None
    if fit is not None:
        scale = int(fit.get("fontScale", "100000")) / 100000.0
        cut = int(fit.get("lnSpcReduction", "0")) / 100000.0
    paras = read_paragraphs(shape, prs, theme, scale, cut)
    if not any(not p.empty for p in paras):
        return []
    x, y, w, h = rect_emu
    width_pt = max(1.0, (w - l_in - r_in) / EMU_PER_PT)
    if nowrap:
        width_pt = 1e6

    # Groups of paragraphs sharing size and alignment; blank paragraphs
    # only add space and close the current group.
    groups: list[tuple[float, list[Para]]] = []     # (top offset pt, paras)
    cursor = 0.0
    current: list[Para] = []
    first = True
    for para in paras:
        if para.empty:
            if current:
                groups.append((start, current))
                current = []
            cursor += para.space_before + line_height(para) + para.space_after
            first = False
            continue
        if current and (abs(para.size - current[-1].size) > 0.5
                        or para.align != current[-1].align):
            groups.append((start, current))
            current = []
        if not current:
            start = cursor + (0 if first else para.space_before)
        cursor += (0 if first else para.space_before) \
            + para_height(para, width_pt) + para.space_after
        current.append(para)
        first = False
    if current:
        groups.append((start, current))
    total = cursor - (paras[-1].space_after if paras else 0)

    box_h_pt = (h - t_in - b_in) / EMU_PER_PT
    offset = 0.0
    if anchor == "ctr":
        offset = (box_h_pt - total) / 2
    elif anchor == "b":
        offset = box_h_pt - total

    out = []
    pw, ph = slide_w / EMU_PER_PT, slide_h / EMU_PER_PT
    for top, group in groups:
        colour = _dominant_color(group)
        size = group[0].size
        height = sum(para_height(p, width_pt) for p in group)
        bx = (x + l_in) / EMU_PER_PT
        by = (y + t_in) / EMU_PER_PT + offset + top
        bw = width_pt if not nowrap else max(
            (w - l_in - r_in) / EMU_PER_PT,
            max(sum(len(r.text) * r.size * _AVG_ADVANCE for r in p.runs)
                for p in group) * 1.1)
        out.append(SlideText(
            x=round(bx / pw, 4), y=round(by / ph, 4),
            w=round(min(bw, pw) / pw, 4), h=round(max(height, size) / ph, 4),
            text=group_text(group, colour),
            font_pt=max(4, int(round(size))),
            color=colour.upper() if colour else "#000000",
            align=group[0].align, locked=False))
    return out
