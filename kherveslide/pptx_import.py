"""Import Microsoft PowerPoint (.pptx) decks into the KherveSlide model.

PowerPoint positions every shape absolutely, which maps cleanly onto our
free-positioned (unlocked) objects: each shape's EMU geometry becomes a
0..1 fraction of the slide, and the deck's page size is set to the
PowerPoint slide size so fonts (in points) land at the right scale.

Text follows PowerPoint's own style inheritance and layout (see
``pptx_text``); pictures keep their crop and effects; groups are
flattened through their child coordinate space; autoshapes become
shapes, connectors and freeform arrows become lines (with arrowheads);
white fade overlays become picture fades. The deck gets a Calibri-metric
typeface and plain "•" bullets. Anything unrecognised is skipped rather
than aborting the import. Needs the optional ``python-pptx`` package.
"""
from __future__ import annotations

from pathlib import Path

from . import pptx_text
from .model import (
    Deck, Slide, SlidePicture, SlideShape, SlideTable, SlideLine,
    SlideVideo, ThemeSpec,
)

_EMU_PER_CM = 360000.0


def available() -> bool:
    try:
        import pptx  # noqa: F401
        return True
    except Exception:
        return False


# Formats XeTeX's \\includegraphics reads; anything else (gif, bmp,
# tiff, wmf, emf...) is read as TeX source and fills the log with garbage.
TEX_IMAGE_EXTS = {"png", "jpg", "jpeg", "pdf", "eps"}


def _save_image(blob: bytes, ext: str, stem: Path) -> str:
    """Write a picture into the media folder in a format XeTeX can include,
    converting with Pillow when needed. Returns its path, or "" when it
    cannot be converted (e.g. WMF/EMF off Windows) — the picture then
    imports as an empty placeholder the user can fill."""
    ext = (ext or "").lower()
    if ext in TEX_IMAGE_EXTS:
        out = stem.with_suffix(f".{ext}")
        out.write_bytes(blob)
        return str(out)
    try:
        import io
        from PIL import Image
        with Image.open(io.BytesIO(blob)) as im:
            out = stem.with_suffix(".png")
            im.convert("RGBA").save(out)
            return str(out)
    except Exception:
        return ""


def _extract_movie(shape, media: Path, n: int) -> tuple[str, str]:
    """Pull the video (and its poster frame) out of a movie shape.
    Returns ``(video_path, poster_path)``, empty strings when missing.

    An embedded video lives as a media part inside the .pptx zip, pointed
    at by the shape's ``<a:videoFile r:link="rIdN"/>``; its bytes are
    written out next to the extracted images. A *linked* (not embedded)
    video keeps its original location. Audio media has ``a:audioFile``
    instead and is skipped."""
    from pptx.oxml.ns import qn

    els = shape._element.xpath(".//a:videoFile")
    if not els:
        return "", ""
    rel = shape.part.rels[els[0].get(qn("r:link"))]
    if rel.is_external:
        video = rel.target_ref
    else:
        part = rel.target_part
        ext = Path(str(part.partname)).suffix or ".mp4"
        media.mkdir(parents=True, exist_ok=True)
        out = media / f"video_{n:03d}{ext}"
        out.write_bytes(part.blob)
        video = str(out)
    poster = ""
    try:
        img = shape.poster_frame
        if img is not None:
            media.mkdir(parents=True, exist_ok=True)
            p = media / f"video_{n:03d}_poster.{img.ext}"
            p.write_bytes(img.blob)
            poster = str(p)
    except Exception:
        poster = ""              # no poster is fine — placeholder shows
    return video, poster


# ------------------------------------------------------- picture effects
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_A14 = "{http://schemas.microsoft.com/office/drawing/2010/main}"


def _pct(el, attr, default=0.0) -> float:
    """A DrawingML percentage attribute (100000 = 100 %) as a fraction."""
    v = el.get(attr) if el is not None else None
    return int(v) / 100000.0 if v not in (None, "") else default


def _clr_hex(parent) -> str:
    """First srgbClr under *parent* as #rrggbb ("" if none / scheme)."""
    if parent is None:
        return ""
    el = parent.find(f".//{_A}srgbClr")
    return f"#{el.get('val').lower()}" if el is not None else ""


def _picture_effects(shape, pic: SlidePicture) -> None:
    """Carry PowerPoint's picture formatting onto *pic*: crop, transparency,
    corrections, colour, artistic effects, soft edges, glow, reflection
    and shadow. Anything unreadable is left at its neutral value."""
    for side in ("left", "top", "right", "bottom"):
        try:
            v = float(getattr(shape, f"crop_{side}"))
        except Exception:
            v = 0.0
        setattr(pic, f"crop_{side[0]}", max(0.0, min(0.9, v)))
    el = shape._element
    geom = el.find(f".//{_A}prstGeom")
    pic.mask = {"ellipse": "ellipse", "roundRect": "rounded",
                "flowChartConnector": "ellipse"}.get(
        geom.get("prst") if geom is not None else "", "")
    blip = el.find(f".//{_A}blip")
    if blip is not None:
        amf = blip.find(f"{_A}alphaModFix")
        if amf is not None:
            pic.opacity = _pct(amf, "amt", 1.0)
        if blip.find(f"{_A}grayscl") is not None:
            pic.recolor = "grayscale"
        duo = blip.find(f"{_A}duotone")
        if duo is not None:
            pic.recolor = "duotone"
            pic.recolor_color = _clr_hex(duo) or "#1f4e79"
        if blip.find(f"{_A}biLevel") is not None:
            pic.recolor = "bw"
        lum = blip.find(f"{_A}lum")
        if lum is not None:
            pic.brightness = _pct(lum, "bright")
            pic.contrast = _pct(lum, "contrast")
        # Office 2010 image effects live in an extension list.
        for eff in blip.iter(f"{_A14}imgEffect"):
            for child in eff:
                tag = child.tag.replace(_A14, "")
                if tag == "brightnessContrast":
                    pic.brightness = _pct(child, "bright")
                    pic.contrast = _pct(child, "contrast")
                elif tag == "saturation":
                    pic.saturation = _pct(child, "sat", 1.0)
                elif tag == "colorTemperature":
                    t = int(child.get("colorTemp", "6500"))
                    pic.temperature = max(-1.0, min(1.0, (6500 - t) / 4000))
                elif tag == "sharpenSoften":
                    pic.sharpness = _pct(child, "amount")
                elif tag.startswith("artistic"):
                    pic.artistic = _ARTISTIC_MAP.get(tag, "")
    effects = el.find(f".//{_A}effectLst")
    if effects is not None:
        ext = el.find(f".//{_A}xfrm/{_A}ext")
        side = min(int(ext.get("cx")), int(ext.get("cy"))) if ext is not None else 0
        soft = effects.find(f"{_A}softEdge")
        if soft is not None and side:
            pic.soft_edge = min(0.5, int(soft.get("rad", "0")) / side)
        glow = effects.find(f"{_A}glow")
        if glow is not None and side:
            pic.glow_color = _clr_hex(glow) or "#ffd966"
            pic.glow_size = min(0.3, int(glow.get("rad", "0")) / side)
        refl = effects.find(f"{_A}reflection")
        if refl is not None:
            pic.reflection = max(0.05, _pct(refl, "endPos", 0.5))
        if effects.find(f"{_A}outerShdw") is not None:
            pic.shadow = True


_ARTISTIC_MAP = {
    "artisticBlur": "blur", "artisticPencilSketch": "pencil",
    "artisticPencilGrayscale": "pencil", "artisticLineDrawing": "line_drawing",
    "artisticMosaicBubbles": "mosaic", "artisticPastelsSmooth": "watercolor",
    "artisticWatercolorSponge": "watercolor", "artisticPaintStrokes":
    "watercolor", "artisticCutout": "posterize", "artisticPhotocopy": "bw",
    "artisticGlowEdges": "glow_edges", "artisticGlowDiffused": "blur",
    "artisticChalkSketch": "pencil", "artisticMarker": "posterize",
    "artisticCement": "emboss", "artisticTexturizer": "emboss",
    "artisticPlasticWrap": "emboss", "artisticFilmGrain": "",
    "artisticCrisscrossEtching": "pencil", "artisticLightScreen": "washout",
}


def _white_fade_overlay(shape):
    """A rectangle filled with a white (background) gradient that runs
    from see-through to solid — PowerPoint's usual way of fading a photo
    into the slide. Returns ``(angle_deg, [(pos, alpha), ...])`` or None."""
    el = shape._element
    if not el.tag.endswith("}sp"):
        return None
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        return None
    grad = el.find(f".//{_A}gradFill")
    if grad is None:
        return None
    stops = []
    for gs in grad.iter(f"{_A}gs"):
        clr = gs[0] if len(gs) else None
        if clr is None:
            return None
        tag = clr.tag.replace(_A, "")
        val = (clr.get("val") or "").lower()
        if not ((tag == "schemeClr" and val in ("bg1", "lt1"))
                or (tag == "srgbClr" and val == "ffffff")
                or (tag == "prstClr" and val == "white")):
            return None
        a = clr.find(f"{_A}alpha")
        stops.append((int(gs.get("pos", "0")) / 100000.0,
                      _pct(a, "val", 1.0) if a is not None else 1.0))
    if len(stops) < 2 or min(a for _, a in stops) > 0.05:
        return None
    if grad.find(f"{_A}path") is not None:
        return "radial", sorted(stops)
    lin = grad.find(f"{_A}lin")
    ang = int(lin.get("ang", "0")) / 60000.0 if lin is not None else 0.0
    return ang % 360, sorted(stops)


def _apply_fade_overlay(objs, rect, overlay) -> None:
    """Turn a white fade overlay into a fade on the picture beneath it.
    The picture is see-through where the overlay is solid, so the slide's
    white shows through exactly as in PowerPoint."""
    ang, stops = overlay
    rx, ry, rw, rh = rect
    pic = next((o for o in reversed(objs) if isinstance(o, SlidePicture)
                and o.path and o.x < rx + rw and rx < o.x + o.w
                and o.y < ry + rh and ry < o.y + o.h), None)
    if pic is None:
        return
    clear = next(p for p, a in stops if a <= 0.05)
    # The solid stop nearest the clear one bounds the fade.
    solid = min((p for p, a in stops if a >= 0.95),
                key=lambda p: abs(p - clear), default=1.0)
    if ang == "radial":
        # Path gradients run from the centre (pos 0) to the edges.
        pic.fade = "radial"
        pic.fade_start = max(0.0, min(1.0, 1 - solid))
        pic.fade_end = max(0.0, min(1.0, 1 - clear))
        return
    axis = int(round(ang / 90.0)) % 4       # 0 →, 1 ↓, 2 ←, 3 ↑
    if axis in (0, 2):
        lo, size, plo, psize = rx, rw, pic.x, pic.w
    else:
        lo, size, plo, psize = ry, rh, pic.y, pic.h

    def at(p):                 # slide coordinate of gradient position p
        return lo + p * size if axis in (0, 1) else lo + size - p * size

    sv, cv = at(solid), at(clear)
    if sv < cv:                # see-through towards the left / top
        pic.fade = "left" if axis in (0, 2) else "top"
        to_t = lambda v: (v - plo) / psize
    else:
        pic.fade = "right" if axis in (0, 2) else "bottom"
        to_t = lambda v: (plo + psize - v) / psize
    pic.fade_start = max(0.0, min(1.0, to_t(sv)))
    pic.fade_end = max(0.0, min(1.0, to_t(cv)))


# ----------------------------------------------------------------- shapes
_PRST_SHAPES = {
    "rect": "rect", "roundRect": "rounded_rect", "snipRoundRect":
    "rounded_rect", "round2SameRect": "rounded_rect", "ellipse": "ellipse",
    "triangle": "triangle", "rtTriangle": "right_triangle",
    "diamond": "diamond", "parallelogram": "parallelogram",
    "trapezoid": "trapezoid", "pentagon": "pentagon_arrow",
    "homePlate": "pentagon_arrow", "hexagon": "hexagon",
    "heptagon": "heptagon", "octagon": "octagon",
    "rightArrow": "arrow_right", "leftArrow": "arrow_left",
    "upArrow": "arrow_up", "downArrow": "arrow_down",
    "leftRightArrow": "double_arrow", "chevron": "chevron",
    "plus": "plus", "mathPlus": "plus", "lightningBolt": "lightning",
    "wedgeRectCallout": "speech", "wedgeRoundRectCallout": "speech",
    "flowChartProcess": "rect", "flowChartAlternateProcess":
    "rounded_rect", "flowChartDecision": "diamond",
    "flowChartTerminator": "rounded_rect", "flowChartConnector": "ellipse",
}
_LINE_PRST = {"line", "straightConnector1", "bentConnector2",
              "bentConnector3", "bentConnector4", "curvedConnector2",
              "curvedConnector3", "curvedConnector4"}
_DASH = {"dash": "dashed", "lgDash": "dashed", "sysDash": "dashed",
         "dashDot": "dashed", "lgDashDot": "dashed", "sysDashDot": "dashed",
         "dot": "dotted", "sysDot": "dotted"}
# Theme line-style widths by lnRef index (Office default theme).
_LNREF_WIDTH = {1: 0.75, 2: 1.0, 3: 1.5}


class _Xform:
    """Maps a shape's EMU rectangle into slide EMU through any groups."""

    def __init__(self, ox=0.0, oy=0.0, sx=1.0, sy=1.0):
        self.ox, self.oy, self.sx, self.sy = ox, oy, sx, sy

    def rect(self, x, y, w, h):
        return (self.ox + x * self.sx, self.oy + y * self.sy,
                w * self.sx, h * self.sy)

    def into_group(self, grp) -> "_Xform":
        x = grp._element.find(f"{_P}grpSpPr/{_A}xfrm")
        if x is None:
            return self
        off, ext = x.find(f"{_A}off"), x.find(f"{_A}ext")
        coff, cext = x.find(f"{_A}chOff"), x.find(f"{_A}chExt")
        if None in (off, ext, coff, cext):
            return self
        cx, cy = int(cext.get("cx")) or 1, int(cext.get("cy")) or 1
        sx = int(ext.get("cx")) / cx
        sy = int(ext.get("cy")) / cy
        gx, gy, _, _ = self.rect(int(off.get("x")), int(off.get("y")), 0, 0)
        return _Xform(gx - int(coff.get("x")) * sx * self.sx,
                      gy - int(coff.get("y")) * sy * self.sy,
                      sx * self.sx, sy * self.sy)


def _xfrm(el):
    x = el.find(f".//{_A}xfrm")
    if x is None:
        return None
    off, ext = x.find(f"{_A}off"), x.find(f"{_A}ext")
    if off is None or ext is None:
        return None
    return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")),
            int(ext.get("cy")), int(x.get("rot", "0")) / 60000.0,
            x.get("flipH") == "1", x.get("flipV") == "1")


def _style_ref(el, name, theme):
    st = el.find(f"{_P}style")
    if st is None:
        return 0, ""
    ref = st.find(f"{_A}{name}")
    if ref is None:
        return 0, ""
    idx = int(ref.get("idx", "0"))
    colour = pptx_text.color_of(ref[0], theme)[0] if len(ref) else ""
    return idx, colour


def _fill(el, theme):
    """(colour, colour2, gradient_dir, alpha) of a shape's fill."""
    if el.get("useBgFill") == "1":           # painted with the background
        return "#" + theme.colors.get("lt1", "FFFFFF"), "", "", 1.0
    sp_pr = el.find(f"{_P}spPr")
    if sp_pr is not None:
        if sp_pr.find(f"{_A}noFill") is not None:
            return "", "", "", 1.0
        sf = sp_pr.find(f"{_A}solidFill")
        if sf is not None and len(sf):
            c, a = pptx_text.color_of(sf[0], theme)
            return c, "", "", a
        gf = sp_pr.find(f"{_A}gradFill")
        if gf is not None:
            stops = sorted(
                (int(gs.get("pos", "0")), pptx_text.color_of(gs[0], theme))
                for gs in gf.iter(f"{_A}gs") if len(gs))
            if stops:
                lin = gf.find(f"{_A}lin")
                ang = int(lin.get("ang", "0")) / 60000.0 if lin is not None \
                    else 90.0
                direction = "horizontal" if int(round(ang / 90)) % 2 == 0 \
                    else "vertical"
                return (stops[0][1][0], stops[-1][1][0], direction,
                        stops[0][1][1])
    idx, colour = _style_ref(el, "fillRef", theme)
    if idx and colour:
        return colour, "", "", 1.0
    return "", "", "", 1.0


def _line(el, theme):
    """(colour, width_pt, style, head, tail) of a shape's outline."""
    ln = el.find(f"{_P}spPr/{_A}ln")
    idx, ref_colour = _style_ref(el, "lnRef", theme)
    colour, width = (ref_colour if idx else ""), _LNREF_WIDTH.get(idx, 0.75)
    style, head, tail = "solid", False, False
    if ln is not None:
        if ln.find(f"{_A}noFill") is not None:
            colour = ""
        sf = ln.find(f"{_A}solidFill")
        if sf is not None and len(sf):
            colour = pptx_text.color_of(sf[0], theme)[0]
        if ln.get("w"):
            width = int(ln.get("w")) / 12700.0
        dash = ln.find(f"{_A}prstDash")
        if dash is not None:
            style = _DASH.get(dash.get("val", ""), "solid")
        he, te = ln.find(f"{_A}headEnd"), ln.find(f"{_A}tailEnd")
        head = he is not None and he.get("type", "none") != "none"
        tail = te is not None and te.get("type", "none") != "none"
    return colour, max(0.25, width), style, head, tail


def _freeform_points(el):
    """A custom-geometry path as cubic Bézier points, fractions of its
    box: start, then (control, control, end) per segment. Straight and
    quadratic pieces are raised to cubics. None if there's no path."""
    path = el.find(f".//{_A}custGeom/{_A}pathLst/{_A}path")
    if path is None:
        return None
    pw, ph = int(path.get("w", "0")) or 1, int(path.get("h", "0")) or 1

    def pt(e):
        return (int(e.get("x")) / pw, int(e.get("y")) / ph)

    pts: list = []
    for cmd in path:
        tag = cmd.tag.replace(_A, "")
        ps = [pt(e) for e in cmd.iter(f"{_A}pt")]
        if tag == "moveTo":
            if pts:
                break                     # first sub-path only
            pts = ps[:1]
        elif not pts:
            continue
        elif tag == "lnTo" and ps:
            a, b = pts[-1], ps[0]
            pts += [a, b, b]
        elif tag == "cubicBezTo" and len(ps) == 3:
            pts += ps
        elif tag == "quadBezTo" and len(ps) == 2:
            a, c, b = pts[-1], ps[0], ps[1]
            pts += [(a[0] + 2 / 3 * (c[0] - a[0]), a[1] + 2 / 3 * (c[1] - a[1])),
                    (b[0] + 2 / 3 * (c[0] - b[0]), b[1] + 2 / 3 * (c[1] - b[1])),
                    b]
    return pts if len(pts) >= 4 else None


def _curve_line(rect, flip_h, flip_v, sw, sh, pts, **kw) -> SlideLine:
    """A SlideLine following freeform Bézier points *pts* (box fractions)."""
    x, y, w, h = rect
    absolute = []
    for fx, fy in pts:
        fx = 1 - fx if flip_h else fx
        fy = 1 - fy if flip_v else fy
        absolute.append(((x + fx * w) / sw, (y + fy * h) / sh))
    (x1, y1), (x2, y2) = absolute[0], absolute[-1]
    lw, lh = x2 - x1, y2 - y1
    curve = []
    if abs(lw) > 1e-4 and abs(lh) > 1e-4:
        for px, py in absolute[1:]:
            curve += [round((px - x1) / lw, 5), round((py - y1) / lh, 5)]
    return SlideLine(x=x1, y=y1, w=lw, h=lh, curve=curve, **kw)


def _segment(x, y, w, h, flip_h, flip_v, sw, sh,
             ends=(0.0, 0.0, 1.0, 1.0)):
    fx1, fy1, fx2, fy2 = ends
    if flip_h:
        fx1, fx2 = 1 - fx1, 1 - fx2
    if flip_v:
        fy1, fy2 = 1 - fy1, 1 - fy2
    x1, y1 = (x + fx1 * w) / sw, (y + fy1 * h) / sh
    x2, y2 = (x + fx2 * w) / sw, (y + fy2 * h) / sh
    return x1, y1, x2 - x1, y2 - y1


def _deck_font(theme) -> str:
    face = (theme.minor_font or "").lower()
    if face.startswith("calibri") or face.startswith("carlito"):
        return "carlito"
    if face in ("arial", "helvetica", "helvetica neue", "arial nova"):
        return "helvetica"
    return ""


def import_pptx(path, media_dir) -> Deck:
    """Parse *path* (a .pptx) into a :class:`Deck`. Extracted images are
    written under *media_dir*."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    sw, sh = prs.slide_width, prs.slide_height
    theme = pptx_text.Theme.from_presentation(prs)
    media = Path(media_dir)
    counters = {"img": 0, "vid": 0}
    slides: list[Slide] = []

    def frac_rect(r):
        x, y, w, h = r
        return (round(x / sw, 4), round(y / sh, 4),
                round(w / sw, 4), round(h / sh, 4))

    def walk(shapes, xf: _Xform, objs: list):
        for shape in shapes:
            try:
                add_shape(shape, xf, objs)
            except Exception:
                continue          # skip a shape we can't read, keep going

    def add_shape(shape, xf, objs):
        el = shape._element
        tag = el.tag.split("}")[1]
        if tag == "grpSp":
            walk(shape.shapes, xf.into_group(shape), objs)
            return
        geo = _xfrm(el)
        if geo is None:
            return
        gx, gy, gw, gh, rot, flip_h, flip_v = geo
        rect = xf.rect(gx, gy, gw, gh)
        x, y, w, h = frac_rect(rect)
        if getattr(shape, "has_table", False):
            tbl = shape.table
            rows = [[pptx_text.escape(c.text) for c in row.cells]
                    for row in tbl.rows]
            objs.append(SlideTable(x=x, y=y, w=w or 0.5, h=h or 0.25,
                                   rows=rows, locked=False))
            return
        st = shape.shape_type
        if st == MSO_SHAPE_TYPE.PICTURE or tag == "pic":
            if st == MSO_SHAPE_TYPE.MEDIA or el.find(
                    f".//{_A}videoFile") is not None:
                counters["vid"] += 1
                video, poster = _extract_movie(shape, media, counters["vid"])
                if video:
                    objs.append(SlideVideo(
                        x=x, y=y, w=w or 0.5, h=h or 0.4,
                        path=video, poster=poster, locked=False))
                return
            img = shape.image
            counters["img"] += 1
            media.mkdir(parents=True, exist_ok=True)
            fname = _save_image(img.blob, img.ext,
                                media / f"img_{counters['img']:03d}")
            pic = SlidePicture(x=x, y=y, w=w or 0.3, h=h or 0.3,
                               path=fname, keep_aspect=False, locked=False,
                               rotation=rot)
            _picture_effects(shape, pic)
            objs.append(pic)
            return
        prst = el.find(f".//{_A}prstGeom")
        prst = prst.get("prst") if prst is not None else ""
        colour, width, style, head, tail = _line(el, theme)
        if tag == "cxnSp" or prst in _LINE_PRST:
            if colour:
                lx, ly, lw, lh = _segment(*rect, flip_h, flip_v, sw, sh)
                objs.append(SlideLine(
                    x=lx, y=ly, w=lw, h=lh, color=colour, width_pt=width,
                    style=style, arrow_start=head, arrow_end=tail,
                    head_size=max(1.0, width / 1.5), locked=False))
            return
        if (overlay := _white_fade_overlay(shape)) is not None:
            _apply_fade_overlay(objs, (x, y, w, h), overlay)
            return
        fill, fill2, grad, alpha = _fill(el, theme)
        curve = _freeform_points(el) if not prst else None
        if curve is not None and not fill:
            if colour:
                objs.append(_curve_line(
                    rect, flip_h, flip_v, sw, sh, curve, color=colour,
                    width_pt=width, style=style, arrow_start=head,
                    arrow_end=tail, head_size=max(1.0, width / 1.5),
                    locked=False))
            return
        if (fill or colour) and (prst in _PRST_SHAPES or curve is not None):
            objs.append(SlideShape(
                x=x, y=y, w=w, h=h, shape=_PRST_SHAPES.get(prst, "rect"),
                fill=fill, fill2=fill2, gradient=grad or "vertical",
                border_color=colour, border_width=width if colour else 1.0,
                style=style, opacity=alpha, rotation=rot, locked=False))
        if getattr(shape, "has_text_frame", False) and \
                shape.text_frame.text.strip():
            objs.extend(pptx_text.text_boxes(shape, prs, theme, rect,
                                             sw, sh))

    for slide in prs.slides:
        objs: list = []
        walk(slide.shapes, _Xform(), objs)
        slides.append(Slide(objects=objs))

    if not slides:
        slides = [Slide()]
    font = _deck_font(theme)
    return Deck(
        slides=slides,
        title=(prs.core_properties.title or "Imported presentation"),
        page_w_cm=round(sw / _EMU_PER_CM, 3) if sw else 0.0,
        page_h_cm=round(sh / _EMU_PER_CM, 3) if sh else 0.0,
        plain_frames=True,
        nav_symbols=False,
        theme_spec=ThemeSpec(enabled=True, font_family=font, bullets="dot"),
        template="PowerPoint import",
    )
