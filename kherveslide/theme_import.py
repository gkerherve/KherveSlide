"""Bring a university (or any) slide template into KherveSlide.

Three ways in, matching what institutions actually hand out:

* ``kit_from_pptx`` — a PowerPoint template (.potx / .pptx): its theme
  colours, fonts, slide-master background, title / footer bars and logo.
* ``kit_from_picture`` — a screenshot or PDF export of one slide: the
  background, the title-bar / footer colours and the main colour are
  picked off the pixels.
* ``install_beamer_theme`` — an existing beamer theme (.sty, or an
  Overleaf / GitHub .zip): installed so ``\\usetheme{<name>}`` works, and
  used as it is.

The first two return a :class:`~kherveslide.theme_kit.ThemeKit`, so the
result opens in the theme wizard for adjusting like any other theme. No Qt
here (pymupdf / numpy / python-pptx only).
"""
from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from .theme_kit import ThemeKit, contrast_text

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_THEME_REL = ("http://schemas.openxmlformats.org/officeDocument/2006/"
              "relationships/theme")

# PowerPoint typeface → serializer.FONT_FAMILIES key ("" = Latin Modern).
_FONT_MAP = (
    (("arial", "helvetica", "liberation sans", "nimbus sans"), "helvetica"),
    (("fira",), "fira"),
    (("source sans",), "sourcesans"),
    (("lato",), "lato"),
    (("open sans",), "opensans"),
    (("roboto",), "roboto"),
    (("palatino", "book antiqua", "pala"), "palatino"),
    (("charter",), "charter"),
    (("times", "georgia", "cambria", "garamond", "minion", "serif"),
     "times"),
)


def assets_dir() -> Path:
    """Where imported logos are kept (next to the saved themes)."""
    from .custom_themes import themes_dir
    d = themes_dir() / "assets"
    d.mkdir(parents=True, exist_ok=True)
    return d


def font_key(typeface: str) -> str:
    low = (typeface or "").lower()
    for names, key in _FONT_MAP:
        if any(n in low for n in names):
            return key
    return ""


def _hex(rgb) -> str:
    return "#" + "".join(f"{int(c):02X}" for c in rgb[:3])


def _luma(hex_color: str) -> float:
    h = hex_color.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _saturation(hex_color: str) -> float:
    h = hex_color.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    mx, mn = max(r, g, b), min(r, g, b)
    return 0.0 if mx == 0 else (mx - mn) / mx


# =============================================================== PowerPoint
def _open_presentation(path: Path):
    """python-pptx refuses .potx/.potm (a different main content type);
    open a copy with the content type switched to a presentation."""
    from pptx import Presentation
    if path.suffix.lower() in (".pptx", ".pptm"):
        return Presentation(str(path))
    tmp = Path(tempfile.mkdtemp(prefix="ks_potx_")) / "template.pptx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml",
                                    b"presentationml.presentation.main+xml")
                data = data.replace(b"presentationml.slideshow.main+xml",
                                    b"presentationml.presentation.main+xml")
            zout.writestr(item, data)
    return Presentation(str(tmp))


def _theme_xml(master):
    from lxml import etree
    for rel in master.part.rels.values():
        if rel.reltype == _THEME_REL:
            return etree.fromstring(rel.target_part.blob)
    return None


def theme_colours(theme) -> dict[str, str]:
    """The theme's colour scheme: dk1, lt1, dk2, lt2, accent1..6."""
    out: dict[str, str] = {}
    if theme is None:
        return out
    scheme = theme.find(f".//{_A}clrScheme")
    if scheme is None:
        return out
    for slot in scheme:
        name = slot.tag.replace(_A, "")
        srgb = slot.find(f"{_A}srgbClr")
        sysc = slot.find(f"{_A}sysClr")
        if srgb is not None:
            out[name] = "#" + srgb.get("val", "000000").upper()
        elif sysc is not None:
            out[name] = "#" + sysc.get("lastClr", "000000").upper()
    return out


def theme_fonts(theme) -> tuple[str, str]:
    """(heading typeface, body typeface) of the theme's font scheme."""
    if theme is None:
        return "", ""
    major = theme.find(f".//{_A}majorFont/{_A}latin")
    minor = theme.find(f".//{_A}minorFont/{_A}latin")
    return ((major.get("typeface", "") if major is not None else ""),
            (minor.get("typeface", "") if minor is not None else ""))


def _shape_fill(shape) -> str:
    try:
        fill = shape.fill
        if fill.type == 1:                     # MSO_FILL.SOLID
            return "#" + str(fill.fore_color.rgb).upper()
    except Exception:
        pass
    return ""


def _background(master) -> str:
    try:
        fill = master.background.fill
        if fill.type == 1:
            return "#" + str(fill.fore_color.rgb).upper()
    except Exception:
        pass
    return ""


def kit_from_pptx(path, assets: Path | None = None) -> ThemeKit:
    """Build a ThemeKit from a PowerPoint template's slide master."""
    path = Path(path)
    prs = _open_presentation(path)
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    master = prs.slide_masters[0]
    theme = _theme_xml(master)
    cols = theme_colours(theme)
    heading_font, body_font = theme_fonts(theme)

    kit = ThemeKit(name=re.sub(r"[_-]+", " ", path.stem).strip()
                   or "Imported theme")
    kit.text = cols.get("dk1", "#000000")
    if kit.text.upper() in ("#000000", "#FFFFFF"):
        kit.text = "#212121" if kit.text.upper() == "#000000" else kit.text
    kit.background = _background(master) or cols.get("lt1", "#FFFFFF")
    # The institution's colour is usually dark 2 or accent 1.
    candidates = [cols.get(k, "") for k in ("dk2", "accent1", "accent2")]
    candidates = [c for c in candidates
                  if c and _saturation(c) > 0.15 and 0.03 < _luma(c) < 0.8]
    kit.primary = candidates[0] if candidates else cols.get("accent1",
                                                            kit.primary)
    accents = [cols.get(k, "") for k in ("accent2", "accent1", "accent3")]
    accents = [c for c in accents if c and c != kit.primary]
    kit.accent = accents[0] if accents else ""
    kit.font = font_key(body_font or heading_font)

    # Bars, rules and the logo live as shapes on the master / its layouts.
    shapes = list(master.shapes)
    for layout in master.slide_layouts:
        shapes += list(layout.shapes)
    top_bar = bottom_bar = ""
    logo_shape = None
    for shp in shapes:
        try:
            x, y, w, h = (int(shp.left or 0), int(shp.top or 0),
                          int(shp.width or 0), int(shp.height or 0))
        except Exception:
            continue
        if not (w and h):
            continue
        if shp.shape_type == 13:                # MSO_SHAPE_TYPE.PICTURE
            area = (w * h) / (sw * sh)
            if area < 0.12 and (logo_shape is None or area <
                                (logo_shape[1])):
                logo_shape = (shp, area, x, y, w, h)
            continue
        colour = _shape_fill(shp)
        if not colour or w < 0.8 * sw:
            continue
        if y + h <= 0.25 * sh and h >= 0.04 * sh:
            top_bar = top_bar or colour
        elif y >= 0.85 * sh:
            bottom_bar = bottom_bar or colour
    if top_bar:
        kit.primary = top_bar
        kit.title_style = "bar"
    else:
        kit.title_style = "plain"
    kit.footer_style = "bar" if bottom_bar else "line"
    kit.title_text = contrast_text(kit.primary)

    if logo_shape is not None and assets is not None:
        shp, _area, x, y, w, h = logo_shape
        try:
            img = shp.image
            assets.mkdir(parents=True, exist_ok=True)
            dest = assets / f"{_safe(kit.name)}_logo.{img.ext}"
            dest.write_bytes(img.blob)
            if img.ext.lower() in ("png", "jpg", "jpeg", "pdf"):
                kit.logo = str(dest)
            else:                    # e.g. emf/svg/wmf: convert to PNG
                kit.logo = _to_png(dest) or ""
            if kit.logo:
                right = x + w / 2 > sw / 2
                top = y + h / 2 < sh / 2
                kit.logo_corner = ("t" if top else "b") + ("r" if right else "l")
                kit.logo_size = round(max(0.05, min(0.3, h / sh)), 3)
        except Exception:
            pass
    return kit


def _safe(name: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in name)[:40] or "theme"


def _to_png(path: Path) -> str | None:
    """Convert a non-LaTeX image (svg, emf…) to PNG with pymupdf if it
    can read it; None otherwise."""
    try:
        import pymupdf
        with pymupdf.open(str(path)) as doc:
            pix = doc[0].get_pixmap(dpi=200, alpha=True)
            out = path.with_suffix(".png")
            pix.save(str(out))
            return str(out)
    except Exception:
        return None


# ================================================================ picture
def _load_rgb(path: Path, width: int = 640):
    """The first page / the image as an (h, w, 3) uint8 array."""
    import numpy as np
    import pymupdf
    if path.suffix.lower() == ".pdf":
        with pymupdf.open(str(path)) as doc:
            page = doc[0]
            z = width / page.rect.width
            pix = page.get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=False)
    else:
        pix = pymupdf.Pixmap(str(path))
        if pix.alpha or pix.n != 3:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix, 0)
    arr = np.frombuffer(pix.samples, dtype=np.uint8)
    arr = arr.reshape(pix.height, pix.stride)[:, :pix.width * 3]
    return arr.reshape(pix.height, pix.width, 3)


def _dominant(pixels, bins: int = 16) -> list[tuple[str, float]]:
    """Most common colours (quantised), as (hex, share) pairs."""
    import numpy as np
    q = (pixels.reshape(-1, 3) // bins).astype(np.int32)
    keys = q[:, 0] * 10000 + q[:, 1] * 100 + q[:, 2]
    uniq, counts = np.unique(keys, return_counts=True)
    order = np.argsort(-counts)
    total = counts.sum()
    out = []
    flat = pixels.reshape(-1, 3)
    for i in order[:12]:
        mask = keys == uniq[i]
        mean = flat[mask].mean(axis=0)
        out.append((_hex(mean.round()), counts[i] / total))
    return out


def _band(arr, y0: float, y1: float) -> str:
    """The colour of a horizontal band if it is one solid colour across
    the whole width (a title / footer bar), else ""."""
    h = arr.shape[0]
    band = arr[int(h * y0):max(int(h * y0) + 1, int(h * y1))]
    top = _dominant(band)
    if top and top[0][1] > 0.6:
        return top[0][0]
    return ""


def kit_from_picture(path) -> ThemeKit:
    """Guess a ThemeKit from an image or PDF of one slide."""
    path = Path(path)
    arr = _load_rgb(path)
    colours = _dominant(arr)
    background = colours[0][0] if colours else "#FFFFFF"
    kit = ThemeKit(name=f"From {path.stem}")
    kit.background = background if _luma(background) < 0.97 else "#FFFFFF"

    def differs(c: str) -> bool:
        h1, h2 = c.lstrip("#"), background.lstrip("#")
        return sum(abs(int(h1[i:i + 2], 16) - int(h2[i:i + 2], 16))
                   for i in (0, 2, 4)) > 60

    top = next((c for c in (_band(arr, 0.0, 0.08), _band(arr, 0.02, 0.12))
                if c and differs(c)), "")
    bottom = next((c for c in (_band(arr, 0.94, 1.0), _band(arr, 0.9, 0.97))
                   if c and differs(c)), "")
    # The main colour: the title bar if there is one, else the most common
    # clearly-coloured pixels.
    vivid = [c for c, share in colours
             if differs(c) and _saturation(c) > 0.25 and share > 0.002]
    if top:
        kit.primary, kit.title_style = top, "bar"
    elif vivid:
        kit.primary, kit.title_style = vivid[0], "plain"
    if bottom:
        kit.footer_style = "bar" if bottom == kit.primary or not vivid \
            else "line"
        if bottom != kit.primary:
            kit.accent = bottom
    else:
        kit.footer_style = "none"
    others = [c for c in vivid if c not in (kit.primary, kit.accent)]
    if not kit.accent and others:
        kit.accent = others[0]
    darks = [c for c, _s in colours if _luma(c) < 0.3 and differs(c)
             and _saturation(c) < 0.35]
    light_bg = _luma(kit.background) > 0.5
    kit.text = darks[0] if darks and light_bg else (
        "#212121" if light_bg else "#F2F2F2")
    kit.title_text = contrast_text(kit.primary)
    return kit


# ============================================================ beamer theme
def install_beamer_theme(path, styles_dir: Path | None = None) -> str:
    """Install a beamer theme from a .sty or a .zip (Overleaf / GitHub
    download) into the user styles folder; returns the name to use with
    \\usetheme. Its images and subfolders are kept with their relative
    layout, so the theme's own \\includegraphics paths still work."""
    from .custom_themes import import_sty
    path = Path(path)
    dest = Path(styles_dir) if styles_dir is not None else None
    if path.suffix.lower() == ".sty":
        name = import_sty(path, dest)
        if not name:
            raise ValueError(
                f"{path.name} is not a beamer theme — the main file is "
                "called beamertheme<Name>.sty")
        return name
    if path.suffix.lower() != ".zip":
        raise ValueError("Choose a beamertheme…sty file or a .zip")
    if dest is None:
        from .style_manager import user_styles_dir
        dest = user_styles_dir()
    tmp = Path(tempfile.mkdtemp(prefix="ks_beamer_"))
    with zipfile.ZipFile(path) as z:
        for member in z.namelist():
            target = (tmp / member).resolve()
            if tmp.resolve() not in target.parents and target != tmp.resolve():
                continue                       # no paths escaping the folder
            z.extract(member, tmp)
    themes = sorted(tmp.rglob("beamertheme*.sty"),
                    key=lambda p: (len(p.parts), p.name))
    if not themes:
        raise ValueError(f"No beamertheme…sty file inside {path.name}")
    root = themes[0].parent                    # the theme's own folder
    dest.mkdir(parents=True, exist_ok=True)
    for f in root.rglob("*"):
        if f.is_dir() or f.name.startswith(".") or "__MACOSX" in f.parts:
            continue
        rel = f.relative_to(root)
        if f.suffix.lower() == ".sty":
            target = dest / f.name            # .sty files go flat
        elif f.suffix.lower() == ".tex":
            continue                          # the example deck, not needed
        else:
            target = dest / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(f, target)
    shutil.rmtree(tmp, ignore_errors=True)
    stem = themes[0].stem
    return stem[len("beamertheme"):]
