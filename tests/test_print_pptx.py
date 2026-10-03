"""Print / print preview (printing.py) and the PowerPoint export
(pptx_export.py). The export's LaTeX compiles are replaced by a fake that
writes a PDF with PyMuPDF, so these run without tectonic."""
import re
from pathlib import Path

import pytest

pymupdf = pytest.importorskip("pymupdf")
pptx = pytest.importorskip("pptx")


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _source_pdf(path: Path, n: int) -> Path:
    doc = pymupdf.open()
    for k in range(n):
        page = doc.new_page(width=453.5, height=255.1)
        page.insert_text((40, 120), f"Slide {k + 1}", fontsize=30)
    doc.save(str(path))
    return path


# ------------------------------------------------------------- printing
def test_parse_range_and_sheets():
    from kherveslide.printing import (PrintSettings, chosen_pages,
                                      parse_range, sheets_for)
    assert parse_range("1-3, 5", 10) == [0, 1, 2, 4]
    assert parse_range("2, 2, 12", 10) == [1]          # no repeats / overflow
    assert parse_range("8-", 10) == [7, 8, 9]
    with pytest.raises(ValueError):
        parse_range("3-1", 10)
    with pytest.raises(ValueError):
        parse_range("abc", 10)
    assert chosen_pages(PrintSettings(which="current", current=4), 10) == [4]
    assert sheets_for(list(range(7)), "handout6") == [[0, 1, 2, 3, 4, 5],
                                                     [6]]


@pytest.mark.parametrize("layout,which,expect", [
    ("slides", "all", 5), ("handout2", "all", 3), ("handout3", "all", 2),
    ("handout6", "all", 1), ("handout9", "all", 1), ("slides", "current", 1),
])
def test_print_to_pdf_lays_out_sheets(qapp, tmp_path, layout, which, expect):
    from PySide6.QtPrintSupport import QPrinter
    from kherveslide.printing import PrintSettings, paint, set_orientation
    doc = pymupdf.open(str(_source_pdf(tmp_path / "s.pdf", 5)))
    out = tmp_path / "out.pdf"
    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(str(out))
    set_orientation(printer, layout)
    s = PrintSettings(layout=layout, which=which, current=2, title="Talk",
                      greyscale=layout == "handout6")
    assert paint(printer, doc, s, dpi=60) == expect
    printed = pymupdf.open(str(out))
    assert printed.page_count == expect
    page = printed[0]
    assert (page.rect.width > page.rect.height) == (layout == "slides")
    text = page.get_text()
    assert "Talk" in text, (                           # the header
        f"text={text!r} fonts={page.get_fonts()} "
        f"drawings={len(page.get_drawings())} "
        f"rawchars={sum(len(sp['chars']) for b in page.get_text('rawdict')['blocks'] for l in b.get('lines', []) for sp in l['spans'])}")


def test_print_dialog_builds_its_preview(qapp, tmp_path):
    from kherveslide.printing import PrintDialog, PrintSettings
    dlg = PrintDialog(_source_pdf(tmp_path / "s.pdf", 4),
                      PrintSettings(title="Talk", current=1))
    dlg.layout_box.setCurrentIndex(dlg.layout_box.findData("handout4"))
    assert "4 slides on 1 page" in dlg.info.text()
    dlg.range_edit.setText("2-3")
    assert dlg.r_range.isChecked() and "2 slides" in dlg.info.text()
    dlg.range_edit.setText("x")
    assert not dlg.print_btn.isEnabled()
    dlg.close()


# ------------------------------------------------------- pptx: text runs
def test_text_paragraphs_runs_and_lists(qapp):
    from kherveslide.pptx_export import needs_picture, text_paragraphs
    paras = text_paragraphs(
        "\\textbf{Goal:} a \\textit{green} LED\n"
        "\\begin{itemize}\n\\item One\n\\begin{enumerate}\n\\item Sub\n"
        "\\end{enumerate}\n\\item Two\n\\end{itemize}")
    assert [p["kind"] for p in paras] == ["p", "ul", "ol", "ul"]
    assert [p["level"] for p in paras] == [0, 0, 1, 0]
    runs = paras[0]["runs"]
    assert runs[0] == ("Goal:", {"bold": True})
    assert any(t == "green" and f.get("italic") for t, f in runs)
    colour = text_paragraphs("\\textcolor[HTML]{C00000}{red} 5\\,\\%")
    assert colour[0]["runs"][0][1]["color"].upper() == "#C00000"
    assert "".join(t for t, _f in colour[0]["runs"]).endswith("5\u2009%")
    assert not needs_picture("Plain $x^2$ text — 20\\,\\%")
    assert needs_picture("$\\frac{a}{b}$")
    assert needs_picture("\\ce{H2O}")


# --------------------------------------------------------- pptx: export
class _FakeCompile:
    """Writes one page per frame; frame titles are drawn as text (a
    \\phantom title is not), and the slide's objects as a marker."""

    def __call__(self, tex, workdir, basename="document", source_dir=None):
        from types import SimpleNamespace
        workdir = Path(workdir)
        workdir.mkdir(parents=True, exist_ok=True)
        frames = tex.split("\\begin{frame}")[1:]
        doc = pymupdf.open()
        for fr in frames:
            page = doc.new_page(width=453.5, height=255.1)
            m = re.search(r"\\frametitle\{(.*)\}", fr)
            if m and not m.group(1).startswith("\\phantom"):
                page.insert_text((10, 20), m.group(1), fontsize=14,
                                 color=(1, 1, 1))
            if "FORMULA" in fr:
                page.insert_text((60, 150), "FORMULA", fontsize=12)
        pdf = workdir / f"{basename}.pdf"
        doc.save(str(pdf))
        return SimpleNamespace(ok=True, pdf_path=str(pdf), log="")


def _deck(tmp_path):
    from PySide6.QtGui import QColor, QImage
    from kherveslide.model import (Deck, Slide, SlideLine, SlidePicture,
                                   SlideShape, SlideTable, SlideText)
    img = QImage(40, 30, QImage.Format_RGB32)
    img.fill(QColor("#3366CC"))
    img.save(str(tmp_path / "pic.png"))
    return Deck(title="Talk", author="Me", slides=[
        Slide(title="Hello", objects=[
            SlideText(x=0.1, y=0.3, w=0.5, h=0.4, font_pt=20,
                      text="\\textbf{Bold} start\n\\begin{itemize}\n"
                           "\\item one\n\\item two\n\\end{itemize}"),
            SlidePicture(x=0.65, y=0.3, w=0.25, h=0.3, path="pic.png",
                         crop_l=0.1),
            SlideTable(x=0.1, y=0.72, w=0.5, h=0.2,
                       rows=[["A", "B"], ["1", "2"]]),
            SlideLine(x=0.1, y=0.25, w=0.4, h=0.0, arrow_end=True),
            SlideShape(shape="star5", fill="#FFCC00", x=0.7, y=0.7),
            SlideText(x=0.1, y=0.55, w=0.3, h=0.1,
                      text="$\\frac{1}{2}$ FORMULA"),
        ]),
        Slide(title="Secret", hidden=True),
        Slide(),
    ])


def test_pptx_export_editable(qapp, tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from kherveslide.pptx_export import ExportOptions, export_pptx
    out = tmp_path / "talk.pptx"
    steps = []
    rep = export_pptx(_deck(tmp_path), out, tmp_path,
                      ExportOptions(dpi=60), compile_fn=_FakeCompile(),
                      progress=lambda m, f: steps.append(f))
    assert steps[-1] == 1.0 and rep.slides == 3
    prs = Presentation(str(out))
    assert prs.core_properties.title == "Talk"
    assert abs(prs.slide_width / prs.slide_height - 453.5 / 255.1) < 0.01
    s1, s2, s3 = prs.slides
    # the theme is the background picture of every slide
    for s in prs.slides:
        assert s._element.cSld.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}bg"
        ) is not None
    # the frame title lands in the title placeholder
    assert s1.shapes.title.text_frame.text == "Hello"
    kinds = [sh.shape_type for sh in s1.shapes]
    assert kinds.count(MSO_SHAPE_TYPE.TABLE) == 1
    assert MSO_SHAPE_TYPE.FREEFORM in kinds                  # the star
    assert MSO_SHAPE_TYPE.LINE in kinds
    pics = [sh for sh in s1.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 2                     # the photo + the equation
    assert abs(pics[0].crop_left - 0.1) < 1e-6
    assert any("as in the PDF" in p.name for p in pics)
    boxes = [sh for sh in s1.shapes if sh.has_text_frame
             and "start" in sh.text_frame.text]
    tf = boxes[0].text_frame
    assert tf.paragraphs[0].runs[0].text == "Bold"
    assert tf.paragraphs[0].runs[0].font.bold
    xml = tf._txBody.xml
    assert xml.count("buChar") == 2 and "one" in tf.text
    table = next(sh for sh in s1.shapes if sh.has_table).table
    assert table.cell(1, 1).text_frame.text == "2"
    line = next(sh for sh in s1.shapes
                if sh.shape_type == MSO_SHAPE_TYPE.LINE)
    assert "tailEnd" in line._element.xml
    # the hidden slide is there, hidden
    assert s2._element.get("show") == "0"
    assert s1._element.get("show") is None
    assert rep.editable >= 5 and rep.pictures == 1


def test_pptx_export_exact_and_without_hidden(qapp, tmp_path):
    from pptx import Presentation
    from kherveslide.pptx_export import (MODE_EXACT, ExportOptions,
                                         export_pptx)
    out = tmp_path / "exact.pptx"
    rep = export_pptx(_deck(tmp_path), out, tmp_path,
                      ExportOptions(mode=MODE_EXACT, include_hidden=False,
                                    dpi=60), compile_fn=_FakeCompile())
    prs = Presentation(str(out))
    assert rep.slides == len(prs.slides) == 2
    assert all(len(s.shapes) == 0 for s in prs.slides)   # pictures only


def test_title_spans_diff_the_blank_title_page():
    from kherveslide.pptx_export import title_spans
    a, b = pymupdf.open(), pymupdf.open()
    for doc, title in ((a, True), (b, False)):
        page = doc.new_page(width=400, height=225)
        page.insert_text((10, 210), "footer 1/3", fontsize=6)
        if title:
            page.insert_text((10, 20), "Results", fontsize=14)
    spans = title_spans(a[0], b[0])
    assert [s[0] for s in spans] == ["Results"]


def test_window_print_and_pptx_menu(qapp, tmp_path, monkeypatch):
    from kherveslide import printing
    from kherveslide.model import Slide
    from kherveslide.window import SlideWindow
    for name in ("_start_compile", "_start_backdrop",
                 "_maybe_autodownload_packages"):
        monkeypatch.setattr(SlideWindow, name, lambda self, *a: None)
    w = SlideWindow()
    labels = [a.text() for a in w.menuBar().actions()[0].menu().actions()]
    assert "Export PowerPoint (.pptx)…" in labels
    assert "Print…" in labels and "Print preview…" in labels
    w.deck.slides = [Slide(title="A"), Slide(title="B", hidden=True),
                     Slide(title="C")]
    w.current = 2
    pdf = _source_pdf(tmp_path / "s.pdf", 2)
    monkeypatch.setattr(SlideWindow, "_slideshow_pdf", lambda self: pdf)
    seen = {}

    def fake_exec(dlg):
        seen["current"] = dlg.s.current
        seen["pages"] = dlg.doc.page_count
        return 0
    monkeypatch.setattr(printing.PrintDialog, "exec", fake_exec)
    w._print()
    assert seen == {"current": 1, "pages": 2}   # C is page 2: B is hidden
    w.close()
