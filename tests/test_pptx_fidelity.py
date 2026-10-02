"""PowerPoint import fidelity: theme colours, style inheritance, layout
into stacked boxes, bullets, groups, autoshapes, connectors and curves."""
import pytest

pptx = pytest.importorskip("pptx")
from pptx import Presentation                                  # noqa: E402
from pptx.dml.color import RGBColor                            # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE          # noqa: E402
from pptx.util import Emu, Inches, Pt                          # noqa: E402

from kherveslide import pptx_import, pptx_text, serializer    # noqa: E402
from kherveslide.model import (                                # noqa: E402
    Deck, Slide, SlideLine, SlideShape, SlidePicture, SlideText,
    deck_from_json, deck_to_json, line_path,
)


def _blank():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _import(prs, tmp_path):
    src = tmp_path / "d.pptx"
    prs.save(str(src))
    return pptx_import.import_pptx(src, tmp_path / "media")


def test_theme_colour_modifiers():
    theme = pptx_text.Theme(colors={"accent1": "4472C4", "lt1": "FFFFFF"})
    from lxml import etree
    el = etree.fromstring(
        '<a:schemeClr xmlns:a="http://schemas.openxmlformats.org/'
        'drawingml/2006/main" val="accent1"><a:lumMod val="50000"/>'
        '</a:schemeClr>')
    colour, alpha = pptx_text.color_of(el, theme)
    assert colour.startswith("#") and colour != "#4472C4" and alpha == 1.0
    el = etree.fromstring(
        '<a:schemeClr xmlns:a="http://schemas.openxmlformats.org/'
        'drawingml/2006/main" val="bg1"><a:alpha val="40000"/>'
        '</a:schemeClr>')
    assert pptx_text.color_of(el, theme) == ("#FFFFFF", 0.4)


def test_sizes_and_blank_paragraphs_split_into_stacked_boxes(tmp_path):
    prs, slide = _blank()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3))
    tf = tb.text_frame
    tf.text = "Heading"
    tf.paragraphs[0].runs[0].font.size = Pt(28)
    for text in ("", "Body one", "Body two"):
        p = tf.add_paragraph()
        p.text = text
        for r in p.runs:
            r.font.size = Pt(14)
    deck = _import(prs, tmp_path)
    boxes = [o for o in deck.slides[0].objects if isinstance(o, SlideText)]
    assert [b.font_pt for b in boxes] == [28, 14]
    assert boxes[1].y > boxes[0].y + boxes[0].h          # below the gap
    assert boxes[1].text == "Body one\nBody two"


def test_bullets_subscripts_and_colours(tmp_path):
    prs, slide = _blank()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    p = tb.text_frame.paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    from lxml import etree
    pPr.append(etree.fromstring(
        '<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/'
        '2006/main" char="&#8226;"/>'))
    r = p.add_run()
    r.text = "N"
    r = p.add_run()
    r.text = "2"
    r.font._rPr.set("baseline", "-25000")
    r = p.add_run()
    r.text = " r"
    r.font.color.rgb = RGBColor(0xFF, 0, 0)
    deck = _import(prs, tmp_path)
    box = next(o for o in deck.slides[0].objects if isinstance(o, SlideText))
    assert box.text.startswith("\\begin{itemize}\n\\item N")
    assert "\\textsubscript{2}" in box.text
    assert "\\textcolor[HTML]{FF0000}{ r}" in box.text
    assert deck.theme_spec.enabled and deck.theme_spec.bullets == "dot"
    assert deck.nav_symbols is False


def test_groups_shapes_and_connectors(tmp_path):
    prs, slide = _blank()
    grp = slide.shapes.add_group_shape()
    oval = grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2), Inches(2),
                                Inches(2), Inches(1))
    oval.fill.solid()
    oval.fill.fore_color.rgb = RGBColor(0x00, 0x80, 0x00)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1),
                                      Inches(5), Inches(6), Inches(6))
    line.line.color.rgb = RGBColor(0, 0, 0xFF)
    ln = line.line._get_or_add_ln()
    from lxml import etree
    ln.append(etree.fromstring(
        '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/'
        '2006/main" type="triangle"/>'))
    deck = _import(prs, tmp_path)
    objs = deck.slides[0].objects
    shape = next(o for o in objs if isinstance(o, SlideShape))
    assert shape.shape == "ellipse" and shape.fill.upper() == "#008000"
    assert abs(shape.x - 2 / 13.333) < 0.01          # group placed right
    seg = next(o for o in objs if isinstance(o, SlideLine))
    assert seg.arrow_end and not seg.arrow_start
    assert seg.color.upper() == "#0000FF"
    serializer.serialize_deck(deck)


def test_use_bg_fill_is_white_not_theme_accent(tmp_path):
    prs, slide = _blank()
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(2),
                                Inches(2))
    sp._element.set("useBgFill", "1")
    deck = _import(prs, tmp_path)
    shape = next(o for o in deck.slides[0].objects
                 if isinstance(o, SlideShape))
    assert shape.fill.upper() == "#FFFFFF"


def test_curved_line_round_trip_and_serialise():
    line = SlideLine(x=0.1, y=0.5, w=0.4, h=-0.2,
                     curve=[0.2, 1.5, 0.6, 2.0, 1.0, 1.0], arrow_end=True)
    back = deck_from_json(deck_to_json(Deck(slides=[Slide(objects=[line])])))
    got = back.slides[0].objects[0]
    assert got.curve == line.curve
    pts = line_path(got)
    assert len(pts) == 4 and pts[-1] == pytest.approx((0.5, 0.3))
    tex = serializer._serialize_line(got, 0.0, 1)
    assert ".. controls" in tex and "TPHorizModule" in tex


def test_freeform_points_become_curve(tmp_path):
    from lxml import etree
    el = etree.fromstring(
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/'
        '2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/'
        '2006/main"><p:spPr><a:custGeom><a:pathLst><a:path w="100" h="100">'
        '<a:moveTo><a:pt x="0" y="100"/></a:moveTo><a:cubicBezTo>'
        '<a:pt x="10" y="0"/><a:pt x="90" y="0"/><a:pt x="100" y="50"/>'
        '</a:cubicBezTo></a:path></a:pathLst></a:custGeom></p:spPr></p:sp>')
    pts = pptx_import._freeform_points(el)
    assert pts == [(0, 1), (0.1, 0), (0.9, 0), (1, 0.5)]
    line = pptx_import._curve_line((0, 0, 100, 100), False, False, 100, 100,
                                   pts, color="#000000")
    assert (line.x, line.y, line.w, line.h) == (0, 1, 1, -0.5)
    assert len(line.curve) == 6


def test_dot_bullets_and_carlito_in_preamble():
    from kherveslide.model import ThemeSpec
    deck = Deck(slides=[Slide()],
                theme_spec=ThemeSpec(enabled=True, font_family="carlito",
                                     bullets="dot"))
    tex = serializer.serialize_deck(deck)
    assert "{carlito}" in tex and "\\textbullet" in tex
    assert serializer.deck_typeface(deck) == "carlito"


def test_picture_oval_mask(tmp_path):
    from PIL import Image
    from kherveslide import image_effects
    path = tmp_path / "a.png"
    Image.new("RGB", (40, 40), "red").save(path)
    pic = SlidePicture(path=str(path), mask="ellipse")
    assert image_effects.has_effects(pic)
    with Image.open(path) as im:
        a = image_effects.render(im, pic).getchannel("A")
    assert a.getpixel((0, 0)) == 0 and a.getpixel((20, 20)) == 255
